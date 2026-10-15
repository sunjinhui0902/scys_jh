"""Update the single V4 seven-sheet workbook from validated, frozen records."""
import collections
import copy
import hashlib
import json
import re
import sys
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote, urlparse

ROOT = Path(__file__).resolve().parents[1]
R = ROOT / '竞品视频分析_2026-10-15'
sys.path.insert(0, str(ROOT / 'scripts'))
sys.path.insert(0, str(R))
from integrate_competitor_supplement import read, atomic, fingerprint, stamp
from phase4_import import excerpt
from phase4_export import put, short, extend, pathlink
from phase2_export import NAMES, table, link, compact
from openpyxl import load_workbook


def atomic_text(path, text):
    path = Path(path)
    tmp = path.with_name(path.name + '.tmp')
    tmp.write_text(text, encoding='utf-8')
    tmp.replace(path)


def main():
    summary = read(R / 'manifest/supplement_01_import_summary.json')
    integrity = read(R / 'manifest/supplement_01_integrity.json')
    records = summary['records']
    out = R / '阅见生活_竞品视频分析_终版.xlsx'
    state_path = R / 'manifest/supplement_01_export_state.json'
    deps = [R / 'manifest/supplement_01_import_summary.json', R / 'manifest/supplement_01_integrity.json',
            R / 'manifest/phase4_templates.json', Path(__file__)]
    digest = hashlib.sha256(''.join(fingerprint(p)['sha256'] for p in deps).encode()).hexdigest()
    if state_path.exists():
        old = read(state_path)
        if old['input_digest'] == digest and old['workbook_sha256'] == fingerprint(out)['sha256']:
            print('EXPORT_RESUME_SKIP: workbook and validated inputs unchanged')
            return
    backup = R / 'manifest/supplement_01_workbook_before.bin'
    if not backup.exists():
        backup.write_bytes(out.read_bytes())
    wb = load_workbook(BytesIO(backup.read_bytes()))
    rows = [json.loads(s) for s in (R / 'manifest/manifest.jsonl').read_text(encoding='utf-8').splitlines()]
    lookup = {x['video_id']: x for x in rows}
    templates = copy.deepcopy(read(R / 'manifest/phase4_templates.json')['templates'])
    for t in templates:
        if t['id'] == 'T01':
            t['refs'] += ['v_2ad33b07cfad615d', 'v_325a2849484e6f15', 'v_b2bbecb4fc6ef0e9', 'v_74a64cf3aca9595d', 'v_d66d5d68290f0e95']
            t['note'] = '新增同账号变体提供阅读/借阅对照；未新增独立模板，不证明表达或转化效果'
        elif t['id'] == 'T02':
            t['refs'].append('v_38eb65b8de747391')
            t['note'] = '可试反差式当代语言开头，须重写；没有完播数据，不采用“大幅降低完播门槛”结论'
        elif t['id'] == 'T03':
            t['refs'].append('v_b199dcc6960a2707')
            t['status'] = '研究结构可用（新增文本补证）'
            t['note'] = '第二参考为80条定向筛选中的贵阳稿，非新增A深度样本；其未核访谈/数字不可作事实。原“第1期”映射仍未采信'
            ss = read(R / 'segments/v_b199dcc6960a2707.aligned.json')['segments']
            proof = {'video_id': 'v_b199dcc6960a2707', 'role': 'Codex全文定向补证，不计入A深度35条',
                     'bridge': excerpt(ss, 512.78, 560.495), 'shopping_entry': excerpt(ss, 591.19, 609.025),
                     'fact_scope': '仅证明稿中提及中国历史名城及橱窗；不证明访谈、地理数字或书中真实收录'}
            t.setdefault('evidence', []).append(proof)
        elif t['id'] == 'T05':
            t['refs'] = ['v_ac516db0b5dced62', 'v_5ddb5a8ad7ccdb86', 'v_849c3889a3d48662']
            t['note'] = '前两条分别作大美中国与365日走遍中国的结构参考，第三条为无商品对照；v_d00cee56eff9e90d实际导向超越百岁，不作旅行图书SKU依据'
    templates.append({'id': 'T06', 'name': '生活问题→解释对照→共读延伸', 'status': 'AI_REVIEW：单样本结构草案',
                      'audience': '推断：有共同阅读需求的家长与孩子；不限定未经证实的年龄/收入层',
                      'books': '实书内容支持的生活知识/亲子共读读物；经济学书名版本仍待核',
                      'seconds': '60～120秒（测试建议）', 'hook': '提出自己设计的具体生活问题',
                      'story': '比较可见现象与解释→区分事实和猜测→共同查证',
                      'bridge': '仅当实书有对应问题/图解时，展示阅读怎样补充解释',
                      'cta': '说明适合的阅读场景，再邀请看目录；不承诺财商、收入或教育效果',
                      'visual': '原创提问卡/自绘示意/获许可实书内容',
                      'topic': '带孩子旅行前，各自写一个想查明白的问题',
                      'refs': ['v_f43b6f1228c21749'],
                      'risk': 'A模板添加的点左下角原话未获转录支持；商业解释、儿童学习效果不能直接采信',
                      'note': '仅一条样本，不凑第二来源；本批不将财商方向扩为主选品路线',
                      'evidence': [{'video_id': 'v_f43b6f1228c21749', 'bridge': records['v_f43b6f1228c21749']['normalized']['bridge']}],
                      'reference_A_template': str(R / 'analysis_a/supplement_01/原创带货模板库_增量.md')})
    template_doc = {'at': stamp(), 'template_count': 6, 'research_usable': 5, 'drafts': 1, 'templates': templates}
    atomic(R / 'manifest/latest_templates.json', template_doc)
    ws1, ws2, ws3 = wb.worksheets[:3]
    extend(ws1, [('全文语义状态（独立任务）', 26), ('最新验收批次', 25)])
    extend(ws2, [('明确商品/下单入口', 24), ('软购买建议（独立口径）', 25)])
    extend(ws3, [('A实际输入核验', 52)])
    for i, row in enumerate(rows, 6):
        vid = row['video_id']
        rec = records.get(vid)
        state = rec['status'] if rec else 'AI_PENDING'
        put(ws1, i, 7, state)
        put(ws1, i, 28, 'AI_PENDING')
        put(ws1, i, 29, '补充01/复核' if vid in summary['records'] and 'supplement_01' in rec.get('record_path', '') else 'Phase4原批' if rec else '未深析')
        row['ai_status'] = state
        row['ai_latest_record'] = rec['record_path'] if rec else None
        row['text_semantic_status'] = 'AI_PENDING'
        put(ws2, i, 34, 'AI_PENDING' if not rec else '有明确入口' if rec['explicit_shopping_entry'] else '转录未见明确入口')
        put(ws2, i, 35, 'AI_PENDING' if not rec else '有' if rec['soft_purchase_suggestion'] else '该核验段未见')
        if not rec:
            continue
        n, original = rec['normalized'], rec['A_original']['fields']
        ev, b, c, vis = rec['record_path'], n['bridge'], n['cta'], n['visual_and_audio']
        for j, value in {22: short(n['theme']), 24: n['category'], 25: rec['input_mode_declared'], 26: rec['acceptance_scope'], 27: link(ev, '当前核验/更正')}.items():
            put(ws1, i, j, value)
        changes = {3: state, 4: n['hook_text'], 5: n['hook_type'], 6: short(n['theme']), 7: short(n['narrative_structure']),
                   8: n['recommended_book'], 9: short(b['text']), 10: b['start_s'] if b['present'] else '无',
                   11: short(c['text']), 12: c['start_s'] if c['present'] else '无', 16: short(n['target_audience']),
                   18: link(rec['source_A_path'], 'A原件'), 20: short(n['product_value'], 150), 21: rec['input_mode_declared'],
                   22: short(original['hook'].get('text_0_3s', '未知')), 23: short(original['hook'].get('full_hook_text', '未提供')),
                   24: b['end_s'] if b['present'] else '无', 25: c['end_s'] if c['present'] else '无', 26: b['type'], 27: c['type'],
                   28: short(n['closing_behavior_A']), 29: '；'.join(n['mentioned_books']) or '无明确书名',
                   30: short(n['product_value'].get('explicit_claims', '未提供')), 31: '推断：' + short(n['product_value'].get('ai_inferred_value', '未提供')),
                   32: link(ev, '校验/原件'), 33: short('；'.join(rec['corrections'] + rec['unresolved_issues']), 180)}
        for j, value in changes.items():
            put(ws2, i, j, value)
        va = original['visual_and_audio']
        changes = {11: vis['visual_type'], 12: vis['subtitle_style'], 13: vis['book_appearance'], 14: vis['audio_style'], 15: state,
                   16: vis['review_method'], 17: rec['input_mode_declared'], 18: link(ev, '画面/核验'),
                   19: short(va.get('visual_type', '') + '；' + va.get('book_appearance', '')), 20: short(va.get('audio_style', '未知')),
                   22: rec.get('A_input_mode_verified', '历史输入声明；按原验收边界')}
        for j, value in changes.items():
            put(ws3, i, j, value)
        ws1.row_dimensions[i].height = 98
        ws2.row_dimensions[i].height = 116
        ws3.row_dimensions[i].height = 94
    ws1['A3'] = '当前：35条代表样本已合并，31条限定范围可用、4条待核；其他297来源仍AI_PENDING；317份全文语义另计。'
    ws2['A3'] = '引文/时间以ASR校验段为准；阅读、借阅、连载与商品入口分列，不推断实际成交。'
    ws3['A3'] = '新增10条40帧由Codex查看；A读取截图过程未证明，声音/完整镜头未知。'
    # Retain source frequencies and historical A descriptions, replace the current account/type summaries.
    old_ws = wb[NAMES[3]]
    extra = []
    for rr in old_ws.iter_rows(min_row=6):
        values = [pathlink(c) for c in rr]
        if values[0] not in ['账号核验汇总', 'Bridge类型', 'CTA类型']:
            extra.append(values + ['不适用'])
    freq = []
    for a in summary['accounts']:
        freq.append(['账号核验汇总', a['account'], a['source_count'], a['unique_count'], a['sample_count'], a['ai_ok'], a['ai_review'],
                     a['book_introduction_count'], a['shopping_entry_count'], '只统计已核验样本，软购买另列', compact(a['sample_ids']),
                     '来源332/SHA317；深度35；入口与成交分开', link(R / 'manifest/supplement_01_import_summary.json', '逐条统计'), a['soft_purchase_count']])
    freq += extra
    for a, values in read(R / 'analysis_a/supplement_01/account_summary.json')['account_distribution_supplement_10'].items():
        freq.append(['A原补充商业分类（未采用）', a, values['new_count'], 10, '原始分类', '不适用', '不适用', '不适用', '不适用',
                     short(values['summary']), '见补充10条', '不把“带货转化为0”或商业率当成交统计', link(R / 'analysis_a/supplement_01/account_summary.json', 'A原汇总'), '不适用'])
    for kind in ['bridge', 'cta']:
        groups = collections.defaultdict(list)
        for vid, rec in records.items():
            if rec['normalized'][kind]['present']:
                groups[rec['normalized'][kind]['type']].append(vid)
        for typ, vids in groups.items():
            freq.append(['Bridge类型' if kind == 'bridge' else 'CTA类型', typ, len(vids), 35, 35, '不适用', '不适用', '不适用', '不适用',
                         '按导入的可定位区段分类', compact(vids), '类型不等于营销效果；未全量重标基线每句行动', link(R / 'manifest/supplement_01_import_summary.json', '证据'), '不适用'])
    wb.remove(old_ws)
    table(wb, NAMES[3], ['类型', '账号/词项/簇', '来源数/频次', '唯一内容/分母', '深度样本', 'AI_OK', 'AI_REVIEW', '引书阅读数',
                        '明确商品入口数', '主题/形式/类型', '证据video_id', '口径与局限', '证据', '软购买建议数'],
          freq, [28, 32, 17, 18, 16, 13, 15, 18, 20, 68, 50, 64, 23, 20], '35条样本：引书17，明确商品入口9，软购买建议1；各类别可重合，均非成交指标。')
    tr = []
    for t in templates:
        tr.append([t['id'], t['status'], t['name'], t['audience'], t['books'], t['seconds'], t['hook'], t['story'], t['bridge'], t['cta'],
                   t['visual'], t['topic'], '；'.join(t['refs']), link(R / 'manifest/latest_templates.json', '当前模板证据'), t['risk'], t['note'], link(t['reference_A_template'], 'A原模板')])
    wb.remove(wb[NAMES[4]])
    ws = table(wb, NAMES[4], ['模板ID', '验收状态', '结构名称', '适用受众（推断）', '书籍范围', '建议时长', '开头方法', '故事方法', '自然引书方法',
                             'CTA方法', '视觉建议', '建议原创选题（待测试）', '参考video_id', '证据', '风险', '使用边界', 'A原模板'],
               tr, [14, 32, 36, 40, 36, 28, 44, 48, 50, 50, 45, 48, 42, 22, 60, 65, 22], '6套结构：5套可作研究参考、1套单样本草案；不是发布成稿，也非效果已验证。')
    for i in range(6, 12):
        ws.row_dimensions[i].height = 145
    old_ws = wb[NAMES[5]]
    ex = [[pathlink(c) for c in rr] for rr in old_ws.iter_rows(min_row=6)
          if rr[2].value not in ['A未决问题', '已整理更正', '第二参考缺映射', '原声音判断边界']]
    for vid, rec in records.items():
        for issue in rec['schema_errors'] + rec['unresolved_issues']:
            ex.append([vid, rec['account'], '当前未决问题', issue, rec['status'], link(rec['record_path'], '当前核验')])
    ex += [['supplement_01', '批次', '交付协议', summary['delivery_protocol_note'], 'PARTIAL（稳定冻结后导入）', link(R / 'manifest/supplement_01_import_summary.json', '批次证据')],
           ['10条', '补充样本', '已更正', '源路径对象、0～3秒、书名、Bridge/CTA时段已逐条更正；A原件保留', '导入版已更正', link(R / 'manifest/supplement_01_import_summary.json', '逐条记录')],
           ['T03', '模板', '补证更新', '贵阳稿已提供第二个真实video_id和文本引书证据；未核实访谈不作为事实', '研究结构可用', link(R / 'manifest/latest_templates.json', '新增文本证据')],
           ['T06', '模板', '单样本草案', '实体书信息和第二来源待核；新增结构暂不列为已充分支持', 'AI_REVIEW', link(R / 'manifest/latest_templates.json', '模板证据')]]
    wb.remove(old_ws)
    table(wb, NAMES[5], ['video_id/模板', '账号', '问题类型', '问题', '状态/处理方式', '证据'], ex,
          [26, 23, 26, 90, 46, 24], '保留原异常，列当前未决项；交付协议、时间和内容验收分别记录。')
    metrics = [
        ['全项目状态', 'PARTIAL', '4条关键待核、A交付协议和真实耗时缺项、全文语义尚未执行'],
        ['更新于', stamp(), '统一北京时间'], ['来源视频', 332, '来源映射全部保留'], ['SHA唯一内容', 317, '15个额外完全重复来源保留'],
        ['本地转录/时间戳', 332, '317份识别及同SHA复用，不代表逐字正确率'], ['基础截图', 1328, '未重做全量截图'],
        ['代表深度样本', 35, '原25＋新10，旧5复核不重复计数'], ['AI_OK（限定范围）', 31, '沿用20＋本批校验11；未证明史实、SKU、声音或效果'],
        ['AI_REVIEW', 4, '2条名人叙事、1条书封/出版社、1条财商图书商品信息'], ['来源AI_PENDING', 297, '332来源减35样本；独立未深析内容282份'],
        ['全量文字语义完成', 0, '独立任务，不能以25/35条深度样本抵扣'], ['全量文字语义待做', 317, '328阅读单元、38批已就绪；当前无A落盘结果'],
        ['定向筛选', 80, '两账号；68纳入、12暂排除，非全量语义AI_OK'], ['原创建议选题', 30, '旅行地理草案；另有10个单点候选，不假定两库无重合'],
        ['已有原创试做稿', 5, '另一路径已生成；0条判为可发布'], ['研究结构/单样本草案', '5 / 1', 'T03新增真实文本来源；T06单样本草案'],
        ['引书/阅读推荐', 17, '35条样本内，是否出现阅读推荐'], ['明确商品/下单入口', 9, '只统计转录中的明确入口或下单路径'], ['软购买建议', 1, '准备一套等购买倾向，未当商品入口'],
        ['A原件不变', '47/47', '原28＋补充19 SHA/mtime保持'], ['源视频不变', integrity['sources_sha_size_mtime_unchanged'], '全量SHA/大小/mtime及目录成员一致'],
        ['数据库', '未打开', '主库/WAL与历史快照一致；SHM修改时间变化，原因未核'], ['本批断点测试', 'PASS', '真实子进程退出75后恢复，3/3检查点不变'],
        ['A分析耗时', '未知', '自报固定时间不能证明真实开始结束'], ['本轮收费API', 0, 'Codex未调用；A既有额度/费用未知'],
        ['当前协作状态', '见协作总控', '完成数量和下一步集中维护；历史报告保留时间口径']]
    wb.remove(wb[NAMES[6]])
    table(wb, NAMES[6], ['指标', '实测/状态', '定义与局限'], metrics, [38, 43, 105], '当前阶段性口径；不得统称317条全部深度分析完成。')
    wb._sheets = [wb[name] for name in NAMES]
    tmp = out.with_name(out.stem + '.tmp.xlsx')
    wb.save(tmp)
    tmp.replace(out)
    # Save current status separately from historical Phase 4 records.
    text = ''.join(json.dumps(r, ensure_ascii=False) + '\n' for r in rows)
    atomic_text(R / 'manifest/manifest.jsonl', text)
    index = R / 'content_library/文案索引.jsonl'
    indexed = [json.loads(s) for s in index.read_text(encoding='utf-8').splitlines()]
    for x in indexed:
        rec = records.get(x['video_id'])
        x['latest_deep_status'] = rec['status'] if rec else 'AI_PENDING'
        x['latest_deep_record'] = rec['record_path'] if rec else None
    atomic_text(index, ''.join(json.dumps(x, ensure_ascii=False) + '\n' for x in indexed))
    ready_path = R / 'handoff/text_light_v1/READY.json'
    ready = read(ready_path)
    archive = R / 'manifest/text_light_ready_before_supplement_01.json'
    if not archive.exists():
        archive.write_bytes(ready_path.read_bytes())
    for ref in ready['files']:
        if Path(ref['path']) == index:
            ref.update(fingerprint(index))
    atomic(ready_path, ready)
    state = {'at': stamp(), 'input_digest': digest, 'workbook_sha256': fingerprint(out)['sha256'], 'workbook': str(out),
             'samples': 35, 'status_counts': summary['combined_status_counts'], 'template_count': 6, 'state': 'STRUCTURAL_DONE_VISUAL_PENDING'}
    atomic(state_path, state)
    print(json.dumps(state, ensure_ascii=True))


if __name__ == '__main__':
    main()
