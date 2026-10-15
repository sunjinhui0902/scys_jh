"""Append the new batch to the one V4 workbook; preserve historical cells."""
from __future__ import annotations
import collections, copy, hashlib, json, re, sys
from pathlib import Path
from urllib.parse import unquote, urlparse
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[1]
R=ROOT/'竞品视频分析_2026-10-15'; BATCH='new_accounts_2026-10-16'; B=R/'increments'/BATCH
sys.path.insert(0,str(R))
from pipeline import atomic_json, sha256, now
from align_segments import token_spans
from phase2_export import safe

def read(p):return json.loads(Path(p).read_text(encoding='utf-8'))
def lines(p):return [json.loads(s) for s in Path(p).read_text(encoding='utf-8').splitlines() if s]
def link(p,label):return {'path':str(p),'label':label}
def brief_ids(values):
    values=list(values);return '；'.join(values[:3])+(f' 等{len(values)}项' if len(values)>3 else '')

def append_row(ws,values,model_row=6,at_row=None):
    i=at_row if at_row is not None else ws.max_row+1
    for j,val in enumerate(values,1):
        src=ws.cell(model_row,j);c=ws.cell(i,j);c._style=copy.copy(src._style)
        c.alignment=copy.copy(src.alignment)
        if isinstance(val,dict) and 'path' in val:
            c.value=val['label'];c.hyperlink=Path(val['path']).as_uri();c.font=copy.copy(src.font)
        else:
            c.value=safe(val)
            if isinstance(val,str):c.number_format='@'
        if c.value=='AI_PENDING':c.fill=PatternFill('solid',fgColor='FFF2CC')
    ws.row_dimensions[i].height=ws.row_dimensions[model_row].height or 48
    return i

def hooks(v):
    raw=read(B/'segments'/f'{v}.json');text=(B/'transcripts'/f'{v}.txt').read_text(encoding='utf-8').strip()
    spans=token_spans(text);stamps=raw.get('raw_timestamp',[])
    if len(spans)!=len(stamps):spans=token_spans(text,keep_apostrophe=True)
    if len(spans)!=len(stamps):return 'AI_PENDING：分块级对齐，不能给逐词0～3秒'
    hits=[i for i,t in enumerate(stamps) if t[0]<3000]
    return text[:spans[hits[-1]][1]] if hits else '0～3秒ASR未识别到词'

def main():
    new=lines(B/'manifest/manifest.jsonl');stats=read(B/'manifest/statistics_summary.json');features=read(B/'clusters/video_features.json')
    cand={x['video_id']:x for x in read(B/'clusters/deep_candidates.json')};issues=read(B/'manifest/local_review_issues.json')
    out=R/'阅见生活_竞品视频分析_终版.xlsx';marker=B/'manifest/workbook_export.json'
    deps=[Path(__file__),B/'clusters/video_features.json',B/'clusters/deep_candidates.json',B/'clusters/current_book_frequencies.json',B/'manifest/local_review_issues.json']
    digest=hashlib.sha256(''.join(sha256(p) for p in deps).encode()).hexdigest()
    expected_current=sha256(out)
    if marker.exists():
        previous=read(marker)
        if expected_current==previous['workbook_sha256'] and previous.get('input_digest')==digest:
            print('WORKBOOK_RESUME_SKIP',flush=True);return
        if expected_current!=previous['workbook_sha256'] or (R/'handoff'/BATCH/'READY.json').exists():
            raise RuntimeError('Workbook or frozen handoff changed; review before updating')
    backup=ROOT/'.local/competitor_review'/BATCH/'before.xlsx';backup.parent.mkdir(parents=True,exist_ok=True)
    if not backup.exists():backup.write_bytes(out.read_bytes())
    wb=load_workbook(backup);original={ws.title:{(c.row,c.column):(c.value,c.hyperlink.target if c.hyperlink else None,c.style_id) for row in ws for c in row if c.value is not None} for ws in wb}
    ws1,ws2,ws3,ws4,ws5,ws6,ws7=wb.worksheets
    assert len(wb.worksheets)==7
    existing={ws1.cell(i,2).value for i in range(6,ws1.max_row+1)}
    assert not existing&{x['video_id'] for x in new},'Duplicate incremental IDs'
    memberships=collections.defaultdict(list)
    for c in read(B/'clusters/current_text_clusters.json'):
        for v in c['new_video_ids']:memberships[v].append(c['cluster_id'])
    for x in new:
        v=x['video_id'];f=features[v];c=cand.get(v);state=read(B/'state'/f'{v}.json');ff=read(B/'frames'/v/'index.json');p='AI_PENDING'
        append_row(ws1,[x['account'],v,Path(x['source_path']).stem[:80],x['duration_s'],'HANDOFF_READY',state['steps']['asr']['status'],p,
            read(B/'manifest/phase2_plan.json')['canonical_by_video_id'][v]['video_id'],'；'.join(f['book_candidates']) or '未命中规则词典',
            f['book_mention_rule'],f['recommendation_rule'],f['shopping_prompt_rule'],'；'.join(memberships[v]) or '无阈值簇',c['rank'] if c else '未选',
            x['size_bytes'],x['sha256'],'未关联数据库',link(x['source_path'],'原视频'),link(B/'transcripts'/f'{v}.txt','全文TXT'),
            link(B/'segments'/f'{v}.aligned.json','时间戳JSON'),link(R/'handoff'/BATCH/'video_inputs'/f'{v}.json','分析输入包'),p,
            c['reason'] if c else '全量文字整理，未入代表候选',p,p,'本地ASR及4静态帧；AI尚未验收',
            link(B/'handoff'/v/'READY.json','本地步骤哈希清单'),p,BATCH])
        values=[p]*ws2.max_column
        vals={1:x['account'],2:v,3:p,4:hooks(v),13:'；'.join(z['text'] for z in f['recommendation_hits'][:2]) or '未命中',
            14:f['recommendation_count'],15:f['shopping_count'],17:link(B/'segments'/f'{v}.aligned.json','本地ASR时间证据'),
            18:'AI_PENDING：尚无已验收A文件',29:'；'.join(f['book_candidates']) or '规则未命中，非无书结论',
            32:link(B/'handoff'/v/'READY.json','输入哈希清单'),33:'词典/规则线索不等于全文语义判断；书名和专名待核'}
        for k,val in vals.items():values[k-1]=val
        append_row(ws2,values)
        values=[x['account'],v]
        for frame in ff:values.extend([frame['time_s'],link(frame['path'],Path(frame['path']).name)])
        values += [p,p,p,p,p,'只完成4静态帧；未分析全程音乐/全部镜头',p,p,p,p,'未补帧','AI_PENDING：未取得本批A输入证明']
        append_row(ws3,values)
    ws4.insert_rows(12,2)
    for offset,account in enumerate(sorted({x['account'] for x in new})):
        xx=[x for x in new if x['account']==account]
        append_row(ws4,['账号（新增本地基础）',account,len(xx),len({x['sha256'] for x in xx}),0,0,0,'AI_PENDING','AI_PENDING',
            'AI_PENDING：全文语义及代表深拆待A',brief_ids(x['video_id'] for x in xx),'规则命中不能当作推荐、商品或转化结论',link(B/'manifest/manifest.jsonl','新增来源清单'),'AI_PENDING'],at_row=12+offset)
    for kind,filename,countkey,denominator in [('当前联合书名候选','current_book_frequencies.json','video_count',stats['combined_unique_contents']),('新增故事/实体线索','story_entity_frequencies.json','new_unique_video_count',len(new))]:
        for f in read(B/'clusters'/filename):
            append_row(ws4,[kind,f['term'],f[countkey],denominator,'AI_PENDING','AI_PENDING','AI_PENDING','AI_PENDING','AI_PENDING',f'旧{f["historical_video_count"]}＋新增{f["new_video_count"]}' if 'historical_video_count' in f else '本地词典出现计数',
                brief_ids(x['video_id'] for x in f['evidence']),'独立内容出现数；书名关系/故事身份尚待AI复核',link(B/'clusters'/filename,'逐条命中证据'),'AI_PENDING'])
    for c in read(B/'clusters/current_text_clusters.json'):
        if not c['incremental']:continue
        append_row(ws4,['增量相似簇',c['cluster_id'],len(c['video_ids']),len(new),'AI_PENDING','AI_PENDING','AI_PENDING','AI_PENDING','AI_PENDING',c['kind'],
            brief_ids(c['video_ids']),'连通分量可传递；共同段可能是CTA，非语义结论',link(B/'clusters/incremental_similarity_pairs.jsonl','成对引句和相似度'),'AI_PENDING'])
    for issue in issues:
        append_row(ws6,[issue['video_id'],issue['account'],issue['type'],issue['issue'],'AI_REVIEW：本地质量待听核',link(issue['evidence'],'对齐证据')])
    mapping={ws7.cell(i,1).value:i for i in range(6,ws7.max_row+1)}
    edits={
        '更新于':(now(),'北京时间；本次仅基础处理及交接准备'),
        '来源视频':(stats['combined_source_paths'],'旧332＋新增83，完整来源映射见current_manifest'),
        'SHA唯一内容':(stats['combined_unique_contents'],'旧317＋新增83；旧15个额外完全重复来源保留'),
        '本地转录/时间戳':(stats['combined_source_paths'],'400份独立识别及同SHA复用；不代表逐字正确率'),
        '基础截图':(stats['combined_source_paths']*4,'旧1328＋新增332帧；不等于完整音画分析'),
        '来源AI_PENDING':(stats['combined_source_paths']-35,'代表深拆35条口径保留；新增83未计AI完成'),
        '全量文字语义待做':(stats['combined_unique_contents'],'旧317待办保留；新增83优先交A，语义验收0'),
        '源视频不变':('见本批不变性报告','新83与旧332按SHA/大小/mtime核对；检查范围和结果另存'),
        '数据库':('未打开','主库指纹保持；WAL/SHM哈希变化，原因未确认；不可宣称辅助文件不变'),
        '本批断点测试':('见本批验收报告','本地分块续跑、就绪包和总表重复运行检查，不能沿用旧PASS')}
    allowed_changes=set()
    ws7.cell(3,1,'当前415来源/400独立内容；基础处理、全文语义与代表深拆分口径计数。')
    allowed_changes.add((ws7.title,3,1))
    for key,(value,note) in edits.items():
        i=mapping[key];ws7.cell(i,2,value);ws7.cell(i,3,note);allowed_changes|={(ws7.title,i,2),(ws7.title,i,3)}
    summaries=[('新增两账号基础处理',len(new),'50＋33；本批AI字段全部AI_PENDING'),
        ('新增代表候选',stats['candidates'],'只说明选样理由，尚未代表深拆验收'),
        ('新增文字轻分析验收',0,'按独立文字任务统计，不能用转录或候选替代')]
    for row in summaries:append_row(ws7,row)
    for ws in [ws1,ws2,ws3]:
        old=ws.cell(3,1).value;ws.cell(3,1,f'当前415来源/400独立内容；35条代表深拆维持31可用、4待核；新增83条AI_PENDING。')
        allowed_changes.add((ws.title,3,1))
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.hyperlink:cell.hyperlink.ref=cell.coordinate
        if ws.tables:
            for table in ws.tables.values():
                table.ref=f'A5:{get_column_letter(ws.max_column)}{ws.max_row}'
                if table.autoFilter:table.autoFilter.ref=table.ref
        ws.auto_filter.ref=f'A5:{get_column_letter(ws.max_column)}{ws.max_row}'
    for title,cells in original.items():
        ws=wb[title]
        for (i,j),(value,href,styleid) in cells.items():
            if (title,i,j) in allowed_changes:continue
            target_row=i+2 if title==ws4.title and i>=12 else i
            assert ws.cell(target_row,j).value==value,(title,i,j,'historical cell value changed')
            assert (ws.cell(target_row,j).hyperlink.target if ws.cell(target_row,j).hyperlink else None)==href
    tmp=out.with_name(out.stem+'.increment.tmp.xlsx');wb.save(tmp)
    check=load_workbook(tmp);assert check.sheetnames==wb.sheetnames and len(check.worksheets)==7
    for title,cells in original.items():
        for (i,j),(value,href,styleid) in cells.items():
            if (title,i,j) in allowed_changes:continue
            target_row=i+2 if title==ws4.title and i>=12 else i
            c=check[title].cell(target_row,j)
            assert c.value==value and (c.hyperlink.target if c.hyperlink else None)==href,(title,i,j,'saved historical value/link changed')
    assert check.worksheets[0].max_row-5==stats['combined_source_paths']
    assert collections.Counter(check.worksheets[0].cell(i,7).value for i in range(6,check.worksheets[0].max_row+1))=={'AI_PENDING':380,'AI_OK':31,'AI_REVIEW':4}
    links=0
    for ws in check:
        for row in ws:
            for c in row:
                assert c.data_type!='f',('unexpected formula',ws.title,c.coordinate)
                if c.hyperlink and c.hyperlink.target.startswith('file:'):
                    parsed=urlparse(c.hyperlink.target);path=unquote(parsed.path).lstrip('/')
                    assert Path(path).exists(),('missing hyperlink',ws.title,c.coordinate,path)
                    links+=1
    assert sha256(out)==expected_current,'Concurrent workbook change before replace'
    tmp.replace(out)
    result={'at':now(),'status':'PASS','workbook_sha256':sha256(out),'prior_workbook_sha256':sha256(backup),
        'input_digest':digest,
        'sheets':check.sheetnames,'source_rows':415,'new_source_rows':83,'AI_PENDING':380,'AI_OK':31,'AI_REVIEW':4,
        'local_hyperlinks':links,'historical_values_and_links_preserved_except_explicit_summary_cells':True,
        'sheet04_two_new_account_rows_inserted_at':12,
        'sheet05_templates_preserved':True,'formula_injection_scan':'PASS','visual_qa':'PENDING'}
    atomic_json(marker,result);print(json.dumps(result,ensure_ascii=False),flush=True)

if __name__=='__main__':main()
