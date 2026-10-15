"""Record independent review decisions without importing rejected A conclusions."""
from __future__ import annotations
import collections
import copy
import json
import pathlib
import re
import shutil
import sys
from urllib.parse import unquote, urlparse
from review_new_competitor_A import ROOT, R, H, A, Q, BATCH, strict, write, sha, now
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

OUT = R / '阅见生活_竞品视频分析_终版.xlsx'
BACKUP = ROOT / '.local/competitor_review' / Q.name / 'before.xlsx'

def workbook():
    evidence = strict(Q / 'evidence_checks.json'); digest = sha(Q / 'evidence_checks.json')
    marker = Q / 'workbook_review_import.json'
    if marker.exists():
        old = strict(marker)
        assert sha(OUT) == old['workbook_sha256'] and old['evidence_sha256'] == digest, 'Workbook or review changed after import'
        print('REVIEW_IMPORT_RESUME_SKIP'); return
    assert len(evidence['results']) == 108 and all(x['review_status'] == 'AI_REVIEW' for x in evidence['results'])
    prior_sha = sha(OUT)
    assert prior_sha == strict(R / 'manifest' / (BATCH + '_job.json'))['workbook_sha256'], 'Unreviewed concurrent workbook edit'
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if BACKUP.exists(): assert sha(BACKUP) == prior_sha
    else: shutil.copy2(OUT, BACKUP)
    wb = load_workbook(BACKUP); assert len(wb.worksheets) == 7
    original = {s.title: {c.coordinate: (c.value, c.hyperlink.target if c.hyperlink else None, copy.copy(c._style))
                         for row in s for c in row if c.value is not None} for s in wb}
    allowed = set(); edits = []
    def setcell(ws, row, col, value, href=None):
        c = ws.cell(row, col); allowed.add((ws.title, c.coordinate))
        if isinstance(value, str) and value.startswith(('=', '+', '-', '@')): value = "'" + value
        c.value = value
        if href is not None: c.hyperlink = pathlib.Path(href).resolve().as_uri()
        if value == 'AI_REVIEW': c.fill = PatternFill('solid', fgColor='FCE4D6')
        edits.append({'sheet': ws.title, 'cell': c.coordinate, 'value': value})
    def append(ws, values):
        row = ws.max_row + 1
        for j, value in enumerate(values, 1):
            c = ws.cell(row, j); model = ws.cell(row - 1, j)
            c._style = copy.copy(model._style)
            setcell(ws, row, j, value)
        ws.row_dimensions[row].height = ws.row_dimensions[row - 1].height or 48
        return row
    byid = collections.defaultdict(dict)
    for item in evidence['results']: byid[item['video_id']][item['kind']] = item
    w1, w2, w3, w4, w5, w6, w7 = wb.worksheets
    for ws in [w1, w2, w3]:
        setcell(ws, 3, 1, '415来源/400独立内容；旧35条深拆保留；新增83份文字、25份代表分析均待复核，验收通过0。')
    counts = collections.Counter()
    for i in range(6, w1.max_row + 1):
        v = w1.cell(i, 2).value
        if v not in byid: continue
        t = byid[v]['text']
        setcell(w1, i, 7, 'AI_REVIEW')
        setcell(w1, i, 22, 'A原件（待复核）', Q / 'frozen' / t['path'])
        setcell(w1, i, 25, 'A自报text_only；本批分析未验收')
        setcell(w1, i, 26, '结构/哈希通过；内容或过程证据未过，详见审核')
        setcell(w1, i, 27, '本批审核证据', Q / 'evidence_checks.json')
        setcell(w1, i, 28, 'AI_REVIEW')
        counts['overview_rows'] += 1
    for i in range(6, w2.max_row + 1):
        v = w2.cell(i, 2).value
        if v not in byid: continue
        setcell(w2, i, 3, 'AI_REVIEW')
        setcell(w2, i, 18, 'A文字原件（待复核）', Q / 'frozen' / byid[v]['text']['path'])
        setcell(w2, i, 21, 'text_only（A自报；未验收）')
        setcell(w2, i, 33, '本批时间、语义概括或真实阅读记录未过；未导入分析结论')
        counts['text_status_rows'] += 1
    for i in range(6, w3.max_row + 1):
        v = w3.cell(i, 2).value
        if v not in byid or 'deep' not in byid[v]: continue
        setcell(w3, i, 15, 'AI_REVIEW')
        setcell(w3, i, 16, '通用帧描述和微动推断未通过；音乐/运镜未知')
        setcell(w3, i, 17, 'A自报text_plus_frames；待核')
        setcell(w3, i, 18, 'A代表原件（待复核）', Q / 'frozen' / byid[v]['deep']['path'])
        setcell(w3, i, 22, '4帧引用已核；具体观察证据不足，未认定逐帧读图完成')
        counts['deep_status_rows'] += 1
    records = {x['video_id']: x for x in strict(H / 'batch_plan.json')['records']}
    for i in range(6, w4.max_row + 1):
        account = w4.cell(i, 2).value
        if account not in {x['account'] for x in records.values()}: continue
        ntext = sum(x['account'] == account for x in records.values())
        ndeep = sum(records[v]['account'] == account and 'deep' in d for v, d in byid.items())
        setcell(w4, i, 1, '账号（新增审阅）')
        for col in [5, 6, 7, 8, 14]: setcell(w4, i, col, 0)
        setcell(w4, i, 10, f'收到文字{ntext}、代表{ndeep}；本批验收0，均AI_REVIEW')
        setcell(w4, i, 12, 'E/F/G仅计已纳入代表样本；本批未纳入，文字另计')
        setcell(w4, i, 13, '本批审核证据', Q / 'evidence_checks.json')
    for v, kinds in sorted(byid.items()):
        text = kinds['text']; codes = {x['code'] for x in text['issues']}
        reason = '文字：等距摘句/未记录分析区间' if 'EVENLY_SAMPLED_SEGMENTS_NOT_SEMANTIC_SUMMARY' in codes else '文字：引用时间不符'
        if 'ANALYSIS_FINISH_AFTER_COMPLETION_MARKER' in codes: reason += '、完成时间矛盾'
        if 'deep' in kinds: reason += '；代表：通用帧描述/微动无据'
        row = append(w6, [v, records[v]['account'], 'A本批分析待复核', reason, 'AI_REVIEW', '逐条审核证据'])
        setcell(w6, row, 6, '逐条审核证据', Q / 'evidence_checks.json')
    keyrows = {w7.cell(i, 1).value: i for i in range(6, w7.max_row + 1)}
    changes = {
        '全项目状态': ('PARTIAL', '新增83文字/25代表未验收；旧35条深拆保留'),
        '更新于': (now(), '北京时间；本次为A结果独立审核'),
        'AI_REVIEW': (4, '仅旧35条代表样本内待核；新增25份代表未纳入'),
        '来源AI_PENDING': (297, '总览另有87个AI_REVIEW来源（旧4＋新增83）和31个限定范围AI_OK'),
        '全量文字语义完成': (0, '本批83份收到，但0通过；不以代表深拆替代'),
        '全量文字语义待做': (400, '未验收400：旧317待做＋新增83待复核'),
        '源视频不变': ('415/415', '本次独立重算SHA/大小/mtime；旧49份证据与原A133份也保持'),
        'A分析耗时': ('无法核实', '78份起止相等，首批3份完成时间矛盾；文件写入用时不能当分析耗时'),
        '当前协作状态': ('仅修订首批5条', '已通过原桌面对话交回，完成后停止；不扩批、不启动旧317'),
        '新增两账号基础处理': (83, '50＋33本地就绪；本批83文字/25代表均待复核'),
        '新增代表候选': (25, '收到25份分析，但0验收纳入；海龟13、De宝12'),
        '新增文字轻分析验收': (0, '83份结构通过、内容未过，全部AI_REVIEW')}
    for key, (value, note) in changes.items():
        row = keyrows[key]; setcell(w7, row, 2, value); setcell(w7, row, 3, note)
    append(w7, ['来源AI_REVIEW（总览）', 87, '旧4＋新增83；不能当已完成深拆数'])
    append(w7, ['本批A文字收到', 83, '50＋33；验收0，待复核83'])
    append(w7, ['本批A代表收到', 25, '13＋12；验收0，待复核25'])
    setcell(w7, 3, 1, '415来源/400内容；文字、代表深拆及本批待复核分别计数。')
    for ws in [w6, w7]:
        ws.auto_filter.ref = f'A5:{get_column_letter(ws.max_column)}{ws.max_row}'
        for table in ws.tables.values():
            table.ref = ws.auto_filter.ref
            if table.autoFilter: table.autoFilter.ref = table.ref
    def unchanged(book):
        checked = 0
        for title, cells in original.items():
            for coord, (value, href, style) in cells.items():
                if (title, coord) in allowed: continue
                c = book[title][coord]
                assert c.value == value and (c.hyperlink.target if c.hyperlink else None) == href and c._style == style, (title, coord)
                checked += 1
        return checked
    preserved = unchanged(wb)
    tmp = OUT.with_name(OUT.stem + '.review.tmp.xlsx'); wb.save(tmp)
    check = load_workbook(tmp)
    assert check.sheetnames == wb.sheetnames and len(check.worksheets) == 7
    assert unchanged(check) == preserved
    assert [tuple(r) for r in check.worksheets[4].values] == [tuple(r) for r in wb.worksheets[4].values]
    source_counts = collections.Counter(check.worksheets[0].cell(i, 7).value for i in range(6, check.worksheets[0].max_row + 1))
    assert source_counts == {'AI_OK': 31, 'AI_REVIEW': 87, 'AI_PENDING': 297}, source_counts
    assert counts == {'overview_rows': 83, 'text_status_rows': 83, 'deep_status_rows': 25}
    hyperlinks = 0
    for ws in check:
        for row in ws:
            for c in row:
                if (ws.title, c.coordinate) in allowed: assert c.data_type != 'f'
                if c.hyperlink:
                    uri = urlparse(c.hyperlink.target)
                    if uri.scheme == 'file':
                        file = unquote(uri.path).lstrip('/')
                        assert pathlib.Path(file).exists(), file
                        hyperlinks += 1
    assert sha(OUT) == prior_sha, 'Workbook changed concurrently'
    tmp.replace(OUT)
    write(marker, {'at': now(), 'status': 'PASS', 'prior_workbook_sha256': prior_sha, 'workbook_sha256': sha(OUT),
                   'evidence_sha256': digest, 'source_rows': 415, 'source_AI_status_counts': dict(source_counts),
                   'text_accepted': 0, 'deep_accepted': 0, 'deep_baseline': 35, 'text_review': 83, 'deep_review': 25,
                   'imported_rejected_analysis_values': 0, 'edits': edits, 'unchanged_cells_verified': preserved,
                   'local_hyperlinks_checked': hyperlinks, 'sheets': check.sheetnames, 'visual_qa': 'PENDING'})
    print(json.dumps({'status': 'PASS', 'counts': dict(counts), 'source_status': dict(source_counts), 'preserved_cells': preserved, 'links': hyperlinks}))

def record():
    marker = strict(Q / 'workbook_review_import.json')
    assert marker['visual_qa'] == 'PASS' and sha(OUT) == marker['workbook_sha256']
    evidence = strict(Q / 'evidence_checks.json'); stage = strict(Q / 'structural_gate.json')
    integrity = strict(Q / 'integrity.json'); handoff = strict(Q / 'revision_handoff.json')
    assert handoff['status'] == 'SENT_AND_WORKING'
    jobpath = R / 'manifest' / (BATCH + '_job.json'); job = strict(jobpath)
    at = job.get('review_completed_at') or now(); event = Q.name
    reportpath = R / 'reports' / 'new_accounts_2026-10-16_0738_A审核报告.md'
    results = collections.defaultdict(dict)
    for item in evidence['results']: results[item['video_id']][item['kind']] = item
    records = {x['video_id']: x for x in strict(H / 'batch_plan.json')['records']}
    decisions = [{'video_id': v, 'account': records[v]['account'], 'text_status': 'AI_REVIEW',
                  'deep_status': 'AI_REVIEW' if 'deep' in x else 'AI_PENDING',
                  'text_accepted': False, 'deep_accepted': False,
                  'text_original': str(Q / 'frozen' / x['text']['path']),
                  'deep_original': str(Q / 'frozen' / x['deep']['path']) if 'deep' in x else None,
                  'issue_codes': sorted({e['code'] for d in x.values() for e in d['issues']})} for v, x in sorted(results.items())]
    write(Q / 'decisions.json', {'reviewed_at': at, 'status': 'PARTIAL', 'accepted_text': 0, 'accepted_deep': 0,
                                'review_text': 83, 'review_deep': 25, 'next_action': 'Only five text revisions; no expansion', 'decisions': decisions})
    decision_byid = {x['video_id']: x for x in decisions}
    def update_jsonl(path, update):
        # Preserve every historical line byte-for-byte; only new batch records change.
        parts = path.read_bytes().splitlines(keepends=True); out = []; updated = 0
        for raw in parts:
            d = json.loads(raw)
            if d.get('video_id') in decision_byid:
                update(d, decision_byid[d['video_id']]); raw = (json.dumps(d, ensure_ascii=False, allow_nan=False)+'\n').encode('utf-8'); updated += 1
            out.append(raw)
        assert updated == 83
        tmp = path.with_name(path.name+'.tmp'); tmp.write_bytes(b''.join(out)); tmp.replace(path)
    def library(d, decision):
        d.update(text_semantic_status='AI_REVIEW', latest_deep_status=decision['deep_status'],
                 semantic_output_path=None, received_semantic_output_path=decision['text_original'],
                 received_deep_output_path=decision['deep_original'], codex_review_record=str(Q / 'decisions.json'),
                 text_semantic_verified=False, representative_deep_verified=False)
    for path in [R / 'content_library/文案索引.jsonl', R / 'content_library' / BATCH / '文案索引.jsonl']: update_jsonl(path, library)
    def manifest(d, decision):
        d.update(ai_status='AI_REVIEW', text_semantic_status='AI_REVIEW', ai_review_record=str(Q / 'decisions.json'),
                 received_A_text_record=decision['text_original'], received_A_deep_record=decision['deep_original'])
    update_jsonl(R / 'manifest/current_manifest.jsonl', manifest)
    statepath = R / 'manifest/collaboration_state.json'; state = strict(statepath)
    state.update(updated_at=at, latest_batch=event, project_status='PARTIAL', pending_source_paths=297,
                 source_paths_outside_integrated_deep_scope=380, source_AI_status_counts=marker['source_AI_status_counts'],
                 workbook_sha256=marker['workbook_sha256'], workbook_validation={k:v for k,v in marker.items() if k!='edits'})
    state['text_semantic'].update(verified_complete=0, pending_unique=317, review_unique=83, unaccepted_unique=400,
                                  units=429, batches=55, A_files_received=83)
    state['text_semantic']['new_batch'].update(pending_unique=0, review_unique=83, unaccepted_unique=83, received=83, verified_complete=0)
    state['current_A_review'] = {'review_id': event, 'reviewed_at': at, 'text_received': 83, 'deep_received': 25,
                                'text_accepted': 0, 'deep_accepted': 0, 'text_review': 83, 'deep_review': 25,
                                'accounts': {account:{'text_received':sum(x['account']==account for x in decisions),
                                                      'deep_received':sum(x['account']==account and x['deep_original'] is not None for x in decisions),
                                                      'text_accepted':0,'deep_accepted':0} for account in sorted({x['account'] for x in decisions})},
                                'report':str(reportpath), 'decisions':str(Q/'decisions.json'), 'can_expand':False,
                                'repair_allowed':True,'repair_scope':'revision_text_01; exactly five text-only outputs',
                                'repair_handoff':handoff, 'review_automation_chain':False}
    state['next_A'] = '只在revision_text_01修订首批5条文字；不继续78条或25条深拆；完成后停下等待验收'
    state['next_Codex'] = '本次定时审核已完成；后续只验5条独立修订最终文件，不自动新建审核链或扩批'
    state['latest_review_integrity'] = integrity
    state['entry_paths']['latest_report'] = str(reportpath)
    state['scheduled_increment'].update(status='REVIEWED_REPAIR_PENDING', counts_not_added_to_completed_totals=False,
                                         review_timing='07:38 review completed; no further automatic review scheduled', review_completed_at=at)
    state['current_local_batch'].update(text_semantic_verified=0, deep_verified=0)
    job.update(updated_at=at, status='REVIEWED_REPAIR_PENDING', review_completed_at=at, can_continue=False,
               continuation_decision='No expansion; five text-only revisions permitted and handed off', received_A_final_files=133,
               received_A_text_files=83, received_A_deep_files=25, verified_text_semantic=0, verified_representative_deep=0,
               reviewed_text_files=83, reviewed_deep_files=25, review_text_status='AI_REVIEW', review_deep_status='AI_REVIEW',
               review_report=str(reportpath), review_decisions=str(Q/'decisions.json'), repair_handoff=handoff,
               repair_scope='Only revision_text_01, five videos; stop after delivery', next_automatic_review=None,
               workbook_sha256=marker['workbook_sha256'])
    write(jobpath, job); write(statepath, state)
    texts = [x for x in evidence['results'] if x['kind']=='text']; deeps = [x for x in evidence['results'] if x['kind']=='deep']
    total_quotes = sum(x['tests']['quote_checks'] for x in evidence['results'])
    matched_quotes = sum(x['tests']['quote_pass'] for x in evidence['results'])
    lines = [
        '# 新两账号｜07:38 A结果独立审核\n',
        f'审核结论记录于：{at}。**全项目PARTIAL。本次收到83份文字＋25份代表分析，验收通过均为0，全部转AI_REVIEW；不扩批，只交回首批5条文字修订。**\n',
        '## 数量与范围\n',
        '|账号|本地有效|收到文字|收到代表分析|文字通过|代表通过|\n|---|---:|---:|---:|---:|---:|\n|失去壳的海龟|50|50|13|0|0|\n|闲不住De宝|33|33|12|0|0|\n',
        '当前415来源/400个SHA独立内容；旧35条已纳入深拆（31限定范围可用、4待核）保留。新增25份只是收到待核，不相加成60条合格深拆。全文语义验收仍0/400，317待做＋83待复核。总览来源状态31 AI_OK／87 AI_REVIEW／297 AI_PENDING，不等同于独立全文或合格深拆数。\n',
        '## 格式、文件与证据门禁\n',
        f'- 23个COMPLETED均合法并通过完成Schema；83文字＋25代表JSON均通过对应Schema。拒绝NaN、Infinity与重复键。所有声明的输入/输出SHA核对，133个最终文件经稳定快照冻结，原件未覆盖。账号/ID/源SHA均吻合，25份代表的源路径字符串吻合。\n- 原文短引句 {matched_quotes}/{total_quotes} 逐字包含于所指段；引用正确仍不能证明结论、CTA分类或完整阅读正确。\n- 83份coverage声明列出的101个单元及首末段ID匹配，但只是声明。未独立证明模型完整阅读，长稿尾段实际理解仍待修订补证。\n- 完成标记、稳定性和SHA经过核查；事后无法观察原子重命名的历史系统调用，不把A的自述当成独立过程证明。\n- account_summary没有随包提供独立Schema，按未验收汇总留存；数目由单视频重新计数。A桌面最终表把两账号深拆13/12写反，JSON和实际文件为海龟13/De宝12。\n',
        '## 未通过的主要原因\n',
        '1. text_02～17共78条story_points为按段号等距抽取的短句，缺乏对各篇论点/故事的概括；reusable_method及原创方向存在批量通用套用。例v_7efe414380952566讲罗布泊与彭加木，方法却套低成本银发旅居；Bridge引用仅“作者说”，不足以证明引书。不能把单元ID勾选当全文理解证据。\n2. 首批5条有13处Hook/Bridge/CTA时间与实际所引段不符，见下表。其余78条起止时间完全相等，未记录正向分析区间；首批3条单视频finished_at晚于本批COMPLETED，过程时间互相矛盾。秒级文件写入区间不能当成模型阅读/分析耗时；不据此推断作者动机。\n3. 25份代表的4帧描述均使用同一通用句式，缺少可对应具体帧的差异观察，并声称“微动”，4张静态帧无法支持。只记录实际自报text_plus_frames，不认定实际逐帧读图已完成。\n4. 引书Bridge与普通转场、推荐与讲解、CTA与生活描述混用；输入文本未见购买CTA不能推出实际无挂车/无商品链接。汇总含截断书名；原创迁移不得凭竞品断言目标实书目录或分册存在。\n',
        '### 首批5条时间差异（秒）\n',
        '|video_id|字段|段ID|A范围|原段范围|\n|---|---|---|---|---|\n']
    for item in texts:
        for e in item['issues']:
            if e['code']=='TIME_NOT_CITED_SEGMENT_RANGE': lines.append(f'|{item["video_id"]}|{e["field"]}|{e["segment_id"]}|{e["claimed"]}|{e["actual"]}|\n')
    lines += ['\n### 静态帧抽查边界\n',
              '本次实际查看3张既有帧：v_fc3786ad66dc77e3的1.000秒为“15分钟/海德格尔/存在与时间”文字版头，1001.034秒为窗边人物插画及书名字幕；v_ad944b26480f2d8b的492.171秒为展台青铜器、字幕“包括多件商周青铜器”。不称已看完25条的100张帧，也不推断连续运动、完整配乐或视频节奏。25份通用描述及微动推断的问题由逐文件字段核对确认。\n',
              '## 不变性、重复导入与成本\n',
              f'- 本次重新计算415条源视频SHA/大小/mtime，415/415保持；旧A/manifest49份、冻结前原A133份、READY的1146个引用均通过。数据库本次未打开、未重新哈希；沿用前轮主库保持/WAL与SHM变化原因未知的未决观察，不宣称全部数据库文件不变。\n- 总表维持7-Sheet，{marker["unchanged_cells_verified"]}个非授权变更单元格的值/样式/链接一致；4667个本地超链接核验。旧35条内容、来源及6套模板保留。只更新83条文字待复核、25条代表待复核及异常证据；不导入任何被拒分析结论。\n- 重复导入、冻结复查及文库历史前缀不变性见review_tests.json；本次不重跑或中断ASR，不把上一轮中断测试冒充本次新做。没有模拟本次导入中断电源。\n- A真实分析耗时无法核实；当前只可报告文件与时间记录的矛盾，不据此估算完整AI成本。无新收费API调用；本次Codex没有原视频/音频外发。A既有会话的额度/计费未独立核验。\n- 6个改动范围预览已查看，原有样式保持；渲染器保存完6图后退出码1，PNG/XLSX另行检查通过。未操作原生Excel验证。预览只在被忽略的本地目录，不产生新恢复包。\n',
              '## 下一步决定\n',
              f'禁止扩批。本次已通过Computer Use在原fuye对话发送独立修订，{handoff["observed_at"]}观察到消息已发出并显示Working。仅修原text_01的5条文字到analysis_a/new_accounts_2026-10-16/revision_text_01/，先真实全文阅读、校时和分类自检，5条交齐后停下等验收。原件冻结，剩余78条/25代表及旧317待办本次不启动。不创建无限定时审核链。\n',
              '[修订任务](../handoff/new_accounts_2026-10-16_revision_01/发给A_仅修订首批5条.md)；[逐条决定](../reviews/new_accounts_2026-10-16_0738/decisions.json)；[证据核对](../reviews/new_accounts_2026-10-16_0738/evidence_checks.json)；[文件不变性](../reviews/new_accounts_2026-10-16_0738/integrity.json)。\n',
              '## 未决ID（83文字，其中25代表）\n', '|video_id|账号|文字|代表|问题代码|\n|---|---|---|---|---|\n']
    for d in decisions: lines.append(f'|{d["video_id"]}|{d["account"]}|AI_REVIEW|{d["deep_status"]}|{", ".join(d["issue_codes"])}|\n')
    reportpath.write_text(''.join(lines), encoding='utf-8', newline='\n')
    current = R / 'reports/最终验收报告.md'; history = R / 'reports/history/最终验收报告_0738审核前_2026-10-16.md'
    if not history.exists(): shutil.copy2(current, history)
    old = current.read_text(encoding='utf-8')
    old = re.sub(r'更新于：.*?\n', f'更新于：{at}。**全项目PARTIAL；本地基础完成。本次A收到83文字/25代表，内容验收均0，全部待复核；只交回5条修订。**\n', old, count=1)
    old = old.replace('380来源AI_PENDING。新增83尚无已验收深拆，不把25个候选计入完成数。', '总览31 AI_OK、87 AI_REVIEW、297 AI_PENDING。新增25份代表未纳入合格深拆，不以交付数冒充验收数。')
    old = old.replace('旧317待办保留，新增83优先交A，按实际完整阅读、证据与Schema核验计数。', '旧317待做，新增83待复核。格式/哈希通过，但语义、校时和输入过程未过；详见本次审核报告。')
    old = re.sub(r'## 当前交接\n.*?\n## 保留', f'## 当前审核与修订\n\n[07:38独立审核报告]({reportpath.name})。不扩批；原对话已收到限定5条文字修订并显示Working，观察于{handoff["observed_at"]}。其余78文字/25代表及旧317不自动启动。\n\n## 保留', old, flags=re.S)
    current.write_text(old, encoding='utf-8', newline='\n')
    # Only the new account records change; source material and all historical A files remain frozen.
    subset = R / 'content_library' / BATCH / '开始使用.md'
    text = subset.read_text(encoding='utf-8')
    text = re.sub(r'更新于：.*?\n', f'更新于：{at}。83份原ASR可查；已收到83文字/25代表分析，验收通过0，均AI_REVIEW。这里只是研究参考库。\n', text, count=1)
    if '## 07:38独立审核' not in text:
        text += '\n## 07:38独立审核\n\n[审核报告](../../reports/'+reportpath.name+')列明时间、等距摘句和通用帧问题。未采用不合格结论生成新模板或原创文案；仅首批5条在独立修订，完整文字原文仍可检索。\n'
    subset.write_text(text, encoding='utf-8', newline='\n')
    entry = R / 'content_library/开始使用.md'; text = entry.read_text(encoding='utf-8')
    text = re.sub(r'更新于：[^。]+。', f'更新于：{at}。', text, count=1)
    text = text.replace('代表候选25条，AI尚未验收。', 'A收到83文字/25代表但均未通过，已标AI_REVIEW；仅交回5条修订。')
    entry.write_text(text, encoding='utf-8', newline='\n')
    control = ROOT / 'docs/竞品分析协作总控.md'; text = control.read_text(encoding='utf-8')
    text = re.sub(r'更新于：[^。]+。', f'更新于：{at}。', text, count=1)
    text = text.replace('旧317待办保留，新增83（101单元/17批）优先，均AI_PENDING', '旧317待做；新83（101单元/17批）已收但均AI_REVIEW，当前只修5条')
    start = text.find('**A（既有待办）**'); end = text.find('\n\n**Codex**', start)
    if start >= 0 and end > start:
        text = text[:start] + '**A（当前限定修订）**：只按new_accounts_2026-10-16_revision_01完成首批5条文字修订到revision_text_01，之后停止等验收。其余78文字、25代表及旧317待办本次不扩批、不启动。' + text[end:]
    text = re.sub(r'（历史计划登记）(?:（历史计划登记）)+', '（历史计划登记）', text)
    block = f'''<!-- NEW_ACCOUNTS_2026_10_16_STATUS_START -->
## 新两账号｜最新审核状态

更新于：{at}。本地83/83就绪、332帧、25候选；415来源/400内容。**全项目PARTIAL：A交付83文字/25代表，内容验收均0，全部待复核。**

- 海龟收到50文字/13代表，De宝33文字/12代表。23完成清单与108单视频Schema及声明SHA通过；78文字等距摘句、首批5条13处校时不符、3条过程时间矛盾，25代表通用帧描述/微动推断不合格。全部ID及证据见[本次审核](../竞品视频分析_2026-10-15/reports/{reportpath.name})。
- 旧35条深拆（31限定范围可用、4待核）与6套模板保留；全文语义0/400，317待做＋83待复核。总览31 AI_OK／87 AI_REVIEW／297 AI_PENDING，与合格深拆口径分开。
- 不扩批。{handoff['observed_at']}已在原fuye对话交回[首批5条修订](../竞品视频分析_2026-10-15/handoff/new_accounts_2026-10-16_revision_01/发给A_仅修订首批5条.md)，界面显示Working。只写revision_text_01，完成后停止。没有新增自动审核链，不启动旧317或其余78/25。
- 本次源视频415/415、旧证据49份及原A133份指纹保持。数据库本次未打开或重新哈希；前轮WAL/SHM变化仍原因未知。游本昌未核逝世叙事及两条钱穆同音问题仍保留，见[原质量清单](../竞品视频分析_2026-10-15/increments/new_accounts_2026-10-16/manifest/local_review_issues.json)。
<!-- NEW_ACCOUNTS_2026_10_16_STATUS_END -->'''
    text = re.sub(r'<!-- NEW_ACCOUNTS_2026_10_16_STATUS_START -->.*?<!-- NEW_ACCOUNTS_2026_10_16_STATUS_END -->', block, text, flags=re.S)
    control.write_text(text, encoding='utf-8', newline='\n')
    runbook = ROOT / 'docs/竞品分析_新账号夜间接续_2026-10-16.md'; text = runbook.read_text(encoding='utf-8')
    block = f'''<!-- NIGHT_RUN_LATEST_START -->
## 最新审核结果

更新于{at}：本地83条/332帧保持。A收到83文字、25代表，Schema/SHA通过，内容验收均0，全部AI_REVIEW；旧35深拆/31可用/4待核保留，全文语义0/400。07:38审核已完成，不重复执行启动任务。

不扩批。原fuye对话已接收首批5条文字独立修订，观察时间{handoff['observed_at']}，状态Working；仅写revision_text_01，修完停止。详见[审核报告](../竞品视频分析_2026-10-15/reports/{reportpath.name})与[修订包](../竞品视频分析_2026-10-15/handoff/new_accounts_2026-10-16_revision_01/READY.json)。本轮不安排新的自动审核链。

基础收尾脚本已加review_completed_at保护，防止把AI_REVIEW重置为AI_PENDING。后续按真实验收适配，不能套旧83/25已完成计数。
<!-- NIGHT_RUN_LATEST_END -->'''
    text = re.sub(r'<!-- NIGHT_RUN_LATEST_START -->.*?<!-- NIGHT_RUN_LATEST_END -->', block, text, flags=re.S)
    runbook.write_text(text, encoding='utf-8', newline='\n')
    readme = R / 'README.md'; text = readme.read_text(encoding='utf-8')
    text = re.sub(r'更新于：[^。]+。', f'更新于：{at}。', text, count=1)
    text = text.replace('332来源／317SHA基础处理完成', '415来源／400SHA基础处理完成')
    text = text.replace('317份全文语义尚未执行，38批已就绪；80条定向筛选不是80条语义完成。', '全文语义验收0/400：旧317待做、新83待复核；本次A83文字/25代表未通过，仅交回5条修订。80条定向筛选不是80条语义完成。')
    text = text.replace('后续A任务：handoff/text_light_v1/，沿用已确认文字输入范围与额度，完整读取317份文字，不新增全量视频拆解。', '当前A只修handoff/new_accounts_2026-10-16_revision_01中的5条文字，之后停止等验收；旧317及其余78文字/25代表本轮不启动。')
    readme.write_text(text, encoding='utf-8', newline='\n')
    log = R / 'manifest/analysis_batches.jsonl'; entries = log.read_bytes().splitlines()
    if not any(json.loads(x).get('batch_id') == event for x in entries if x.strip()):
        with log.open('ab') as stream: stream.write((json.dumps({'batch_id':event,'at':at,'status':'PARTIAL','kind':'independent_A_review',
            'text_received':83,'deep_received':25,'text_accepted':0,'deep_accepted':0,'deep_cumulative':35,
            'source_status_counts':marker['source_AI_status_counts'],'review_report':str(reportpath),
            'repair_handoff':handoff,'workbook_sha256':marker['workbook_sha256']},ensure_ascii=False)+'\n').encode('utf-8'))
    journal = ROOT / '05_爆款知识库与复盘/06_竞品分析更新日志.md'; text = journal.read_text(encoding='utf-8')
    if '｜07:38 A独立审核（不扩批）' not in text:
        text += f'\n## {at}｜07:38 A独立审核（不扩批）\n\n收到83文字/25代表；23完成清单、108单视频Schema与SHA通过，但内容均未过，标AI_REVIEW。首批5条13处时间不符，78条等距摘句/无区间，25代表通用帧描述与微动无据。旧35深拆/31可用/4待核、6模板保留，全文语义0/400。已在原fuye对话交回5条文字独立修订，之后停止；不创建自动审核链。[本次报告](../竞品视频分析_2026-10-15/reports/{reportpath.name})。\n'
    journal.write_text(text, encoding='utf-8', newline='\n')
    print('REVIEW_RECORDS_UPDATED', event)

if __name__ == '__main__':
    if '--record' in sys.argv: record()
    else: workbook()
