"""Evidence-backed single-file review and incremental project record."""
from __future__ import annotations
import argparse, collections, copy, datetime as dt, hashlib, json, pathlib, re, shutil, sys
from urllib.parse import unquote, urlparse
ROOT=pathlib.Path(__file__).resolve().parents[1];R=ROOT/'竞品视频分析_2026-10-15';KEY='history_city_2026-10-19';B=R/'increments'/KEY
def now():return dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).isoformat(timespec='seconds')
def read(p):return json.loads(pathlib.Path(p).read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(pathlib.Path(p).read_bytes()).hexdigest()
def write(p,d):
    p=pathlib.Path(p);p.parent.mkdir(parents=True,exist_ok=True);t=p.with_name(p.name+'.tmp');t.write_text(json.dumps(d,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8',newline='\n');t.replace(p)
def source():return read(B/'manifest/source.json')
def evidence(ss,*ids):
    z=[ss[f's{i:06d}'] for i in ids]
    return {'segment_ids':[s['segment_id'] for s in z],'quote':''.join(s['text'] for s in z),'start_s':z[0]['start_ms']/1000,'end_s':z[-1]['end_ms']/1000}

def analysis():
    if (B/'manifest/READY.json').exists():
        for ref in read(B/'manifest/READY.json')['files']:assert sha(ref['path'])==ref['sha256']
        print('SINGLE_ANALYSIS_RESUME_SKIP');return
    x=source();v=x['video_id'];a=read(B/'segments'/f'{v}.aligned.json');ss={s['segment_id']:s for s in a['segments']}
    text=(B/'transcripts'/f'{v}.txt').read_text(encoding='utf-8').strip();assert ''.join(s['text'] for s in ss.values())==text
    frames=read(B/'frames'/v/'index.json')+read(B/'manifest/supplementary_frames.json')
    from PIL import Image
    for f in frames:
        with Image.open(f['path']) as im:im.verify()
    sys.path.insert(0,str(R));from align_segments import token_spans
    raw=read(B/'segments'/f'{v}.json');spans=token_spans(text);stamps=raw['raw_timestamp'];assert len(spans)==len(stamps)
    indices=[i for i,t in enumerate(stamps) if t[1]<=3000];opening=text[:spans[indices[-1]][1]] if indices else ''
    chapters=[]
    for title,start,end in [('续集承接与评论悬念',1,5),('西安案例及与成都对比',6,30),('祖辈带娃的讲解焦虑',31,42),('图书及使用场景',43,70),('价格锚点、购物CTA与售后声称',71,87),('十城回顾与情感收束',88,116),('评论、关注、转发与下期预告',117,127)]:
        chapters.append({'title':title,'first_segment_id':f's{start:06d}','last_segment_id':f's{end:06d}','start_s':ss[f's{start:06d}']['start_ms']/1000,'end_s':ss[f's{end:06d}']['end_ms']/1000})
    book='大美中国：一生要去的101个地方'
    examples={'hook':evidence(ss,1),'bridge':evidence(ss,32,33,34,35,36),'book_entry':evidence(ss,43,44),
              'purchase':evidence(ss,78,79,80),'after_sales_claim':evidence(ss,81,82,83,84,85,86,87),
              'follow':evidence(ss,122,123,124),'next_episode':evidence(ss,127)}
    observations=[(1,'古风城市地图与手指示意，字幕“上一集咱们盘点了”；底部有AI生成与网络下载说明'),(2,'同类地图场景，字幕“长见识的历史名城”'),
        (3,'古风持书男子，字幕“书里都给您准备好了”；未见可对应口播书名的清晰商品封面'),(4,'男子指向古风地图，字幕“中国还有哪些地方”'),
        (5,'陶俑制作场景示意，字幕“秦始皇的陪葬品”；不能用作真实兵马俑形制证据'),(6,'竹椅与茶桌场景，字幕“少一样东西”'),
        (7,'古风成人与儿童、城门背景，字幕“这确实是个问题”'),(8,'成人与儿童共看图册，字幕“下一站去哪儿”'),
        (9,'古风城墙与楼阁示意，字幕“开封”；建筑真实性未核'),(10,'成人与儿童看书，字幕“那些字就不再是字”'),
        (11,'城墙楼阁示意，字幕“明朝修的”；支持原ASR“古朝修的”的字幕级更正'),(12,'山水书封示意；封面文字无法对应口播《大美中国》书名，非真实商品证明'),
        (13,'近景男子，字幕“左下角的小黄车”；本地视频帧不能证明平台挂载仍有效'),(14,'包裹中多本书封示意，文字无法对应口播书名，不能作为实际商品/套数证明')]
    visual=[{**next(f for f in frames if f['index']==i),'sha256':sha(next(f['path'] for f in frames if f['index']==i)),'observation':note} for i,note in observations]
    obj={'schema_version':'codex-single-content-review-1.0','video_id':v,'account':x['account'],'source_video_path':x['source_path'],'source_sha256':x['sha256'],
         'input_mode':'text_plus_frames','analyzed_at':now(),'status':'AI_REVIEW','text_semantic_status':'AI_OK','reviewer':'Codex',
         'coverage':{'read_all_segments':True,'segment_count':127,'first_segment_id':'s000001','last_segment_id':'s000127','missing_ranges':[],
                     'reviewed_frame_count':14,'continuous_video_audio_reviewed':False,'reading_interval':'Not separately clocked; analyzed_at records result creation, not inference duration.'},
         'fields':{'theme':'祖辈带孙辈游历史名城，借不会讲/怕讲错引入旅行人文图书','category':'历史名城/隔代陪伴/场景化图书推荐',
                   'mentioned_books':[book],'recommended_book':book,'book_status':'口播商品名；ISBN、版本、真实商品封面及实际挂载未核',
                   'hook':{**examples['hook'],'opening_0_3_asr':opening,'type':'续集承接＋评论提问；完整首句不冒充0～3秒'},
                   'narrative_structure':'续集承接—西安案例—祖辈讲解焦虑—图书解决方案—价格/购物/售后声称—十城与亲情回顾—互动/下集',
                   'bridge':{**examples['bridge'],'type':'先解释不知道去哪、不会讲、怕误导，再给书；首次书名入口116.710秒'},
                   'cta':{**examples['purchase'],'type':'明确购买口播；有小黄车指引，不代表实际商品链接核验通过'},
                   'visual_and_audio':{'visual':'古风地图、情景人物、祖孙共读、书案/书封示意；黄字黑底字幕；底部AI/网络素材声明',
                                       'audio':'仅本地ASR；声音身份、音乐与连续音画对应未核','camera_movement':'未知；静态帧不能证明运动轨迹'},
                   'book_pitch_seconds':round(228.945-116.710,3),'book_pitch_fraction':(228.945-116.710)/x['duration_s'],
                   'pain_and_pitch_seconds':round(228.945-83.040,3),
                   'method':'从买旅行知识，转为祖辈拥有选地、备课、现场讲故事、回家共读的具体能力；以玩具支出作价格锚点，以孩子成长作时间紧迫感'},
         'chapters':chapters,'evidence':[{**e,'supports':k} for k,e in examples.items()],'visual_evidence':visual,
         'issues':[{'type':'PROMISE_DELIVERY_GAP','detail':'开头承诺接着讲前五，正文详讲集中于西安，随后售书和十城一句话回顾；本文件没有逐一充分展开前五，不据此断定下载缺失','evidence_segments':['s000004','s000006','s000024','s000030','s000090']},
                   {'type':'PRODUCT_VISUAL_MISMATCH','detail':'120秒和220秒书封文字无法对应口播商品，实物书、内页、套数和版本未核','frame_indices':[12,14]},
                   {'type':'UNVERIFIED_PRODUCT_AND_SALES_CLAIMS','detail':'适合孩子的故事讲法、实景照片/手绘地图、先用后付与七天退等均为素材声称，未核实际SKU/活动'},
                   {'type':'ASR_PROPER_NOUNS','detail':'保留原ASR：“古朝修的”在50秒字幕为“明朝修的”；“大宫城”“张一为”“古属文明”等不静默替换，其他候选待核'}],
         'limitations':['全文127段已读，14张静态帧已查看；未逐句听核音频、未观看全部连续画面，字错率未测',
                        '排名、几十万工匠、十三朝及教育效果等作为视频声称，未独立核实史实或效果',
                        '无播放、停留、商品点击或成交数据，不能称已验证爆款；本条仍适合结构研究'],
         'publication_date':'文件名包含20261015，未核真实发布时间'}
    schema=read(R/'handoff/new_accounts_2026-10-16/deep_analysis_schema.json');schema['properties']['account']={'const':x['account']};schema['title']='Single requested-file analysis; account scope extension only'
    write(B/'manifest/analysis_schema.json',schema)
    sys.path.insert(0,str(ROOT/'.local/schema-validation'));from jsonschema import Draft202012Validator,FormatChecker
    Draft202012Validator(schema,format_checker=FormatChecker()).validate(obj)
    for e in obj['evidence']:
        z=[ss[i] for i in e['segment_ids']];assert e['quote']==''.join(s['text'] for s in z) and len(e['quote'])<=200
        assert e['start_s']==z[0]['start_ms']/1000 and e['end_s']==z[-1]['end_ms']/1000
    write(B/'analysis.json',obj)
    current=pathlib.Path(x['source_path']);assert sha(current)==x['sha256'] and current.stat().st_size==x['size_bytes'] and current.stat().st_mtime_ns==x['mtime_ns']
    write(B/'manifest/content_validation.json',{'at':now(),'schema':'PASS','identity':'PASS','exact_quote_time_checks':len(obj['evidence']),'full_text_segments_read':127,'static_frames_read':14,
        'source_unchanged':True,'text_semantic_status':'AI_OK','deep_status':'AI_REVIEW','facts_and_product_verified':False,'original_video_or_audio_uploaded':False,
        'local_processing':read(B/'manifest/local_complete.json'),'maximum_inter_segment_gap_s':max((s['start_ms']-p['end_ms'])/1000 for p,s in zip(a['segments'],a['segments'][1:]))})
    refs=[B/'analysis.json',B/'manifest/analysis_schema.json',B/'manifest/content_validation.json',B/'transcripts'/f'{v}.txt',B/'transcripts'/f'{v}.timed.txt',B/'segments'/f'{v}.aligned.json',*[pathlib.Path(f['path']) for f in frames]]
    write(B/'manifest/READY.json',{'at':now(),'kind':'local_single_file_research','files':[{'path':str(p),'sha256':sha(p),'size_bytes':p.stat().st_size} for p in refs],'A_handoff_requested':False})
    print('SINGLE_REVIEW_VALIDATED',v)

def workbook():
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter
    out=R/'阅见生活_竞品视频分析_终版.xlsx';marker=B/'manifest/workbook_import.json';obj=read(B/'analysis.json');x=source();v=x['video_id']
    if marker.exists():
        assert sha(out)==read(marker)['sha256'];print('SINGLE_WORKBOOK_RESUME_SKIP');return
    assert sha(out)==read(R/'manifest/collaboration_state.json')['workbook_sha256']
    backup=ROOT/'.local/competitor_review'/KEY/'before.xlsx';backup.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(out,backup)
    wb=load_workbook(backup);assert len(wb.worksheets)==7
    originals={ws.title:{c.coordinate:(c.value,c.hyperlink.target if c.hyperlink else None,copy.copy(c._style)) for row in ws for c in row if c.value is not None} for ws in wb}
    allowed=set()
    def put(ws,i,j,val):
        c=ws.cell(i,j);allowed.add((ws.title,c.coordinate))
        if isinstance(val,tuple):c.value=val[1];c.hyperlink=pathlib.Path(val[0]).resolve().as_uri()
        else:c.value=("'"+val) if isinstance(val,str) and val.startswith(('=','+','-','@')) else val
        if c.value=='AI_REVIEW':c.fill=PatternFill('solid',fgColor='FCE4D6')
    def add(ws,values):
        i=ws.max_row+1
        for j,val in enumerate(values,1):ws.cell(i,j)._style=copy.copy(ws.cell(i-1,j)._style);put(ws,i,j,val)
        ws.row_dimensions[i].height=max(ws.row_dimensions[i-1].height or 48,48)
        return i
    w1,w2,w3,w4,w5,w6,w7=wb.worksheets;assert v not in [w1.cell(i,2).value for i in range(6,w1.max_row+1)]
    link=lambda p,label:(p,label)
    report=R/'reports/单条拆解_食力过日子_带孙辈长见识_历史名城下集.md';f=obj['fields'];e={z['supports']:z for z in obj['evidence']}
    add(w1,[x['account'],v,x['source_title'],x['duration_s'],'LOCAL_OK','OK','AI_REVIEW',v,f['recommended_book'],'明确提及','明确推荐/口播商品',
        '明确小黄车口播','新SHA；未重算全库相似文案','用户指定单条',x['size_bytes'],x['sha256'],'未打开数据库',link(x['source_path'],'原视频'),
        link(B/'transcripts'/f'{v}.txt','全文TXT'),link(B/'segments'/f'{v}.aligned.json','时间戳JSON'),link(B/'manifest/READY.json','单条研究包'),
        link(B/'analysis.json','Codex分析'), '用户指定；祖孙名城选题',f['theme'],'text_plus_frames','127段全文＋14静态帧；商品/史实待核',
        link(B/'manifest/content_validation.json','校验记录'),'AI_OK',KEY])
    values=['AI_PENDING']*35
    actual={1:x['account'],2:v,3:'AI_REVIEW',4:f['hook']['opening_0_3_asr'],5:'续集承接/评论提问',6:f['theme'],7:f['narrative_structure'],8:f['recommended_book'],
        9:f['bridge']['quote'],10:f['bridge']['start_s'],11:f['cta']['quote'],12:f['cta']['start_s'],13:e['book_entry']['quote'],14:'未计数',15:'未计数',
        16:'口播商品；版本/内容/售后未核',17:link(B/'segments'/f'{v}.aligned.json','ASR时间证据'),18:link(B/'analysis.json','本条分析'),
        19:'带孙辈不知道去哪/不会讲/怕讲错',20:'选地、备课、现场讲解、回家共读（素材声称）',21:'text_plus_frames',22:f['hook']['opening_0_3_asr'],23:f['hook']['quote'],
        24:f['bridge']['end_s'],25:f['cta']['end_s'],26:'内容困难到图书方案',27:'明确购买；另有评论/关注/转发/下集',28:'见原文s000117—127',
        29:f['recommended_book'],30:'儿童讲法/照片/地图均为素材声称',31:f['method'],32:link(B/'manifest/content_validation.json','证据与身份校验'),
        33:'文字理解通过；前五承诺缺口、书封不符与商品/史实待核',34:'口播小黄车；当前挂载未知',35:'全量文字轻分析1条，非营销效果证明'}
    for j,val in actual.items():values[j-1]=val
    add(w2,values)
    base=read(B/'frames'/v/'index.json');values=[x['account'],v]
    for z in base:values.extend([z['time_s'],link(z['path'],pathlib.Path(z['path']).name)])
    values+=['古风地图/情景人物/祖孙共读','黄字黑底；底部AI/网络素材声明','口播商品与120/220秒封面文字无法对应','声音/配乐/运动未知','AI_REVIEW',
             '4基础帧＋10补帧；非连续视听验收','text_plus_frames',link(B/'analysis.json','14帧逐张观察'),'未知','未知',link(B/'manifest/supplementary_frames.json','补帧10张'),
             'Codex实际阅读127段与14帧；时间与引句校验通过']
    assert len(values)==22;add(w3,values)
    add(w4,['账号（单条指定，非全盘）',x['account'],1,1,1,0,1,1,1,'祖孙讲解困难到旅行人文书',v,'只登记这一条；全库频次仍为历史400内容口径',link(report,'单条报告'),1])
    for issue in obj['issues']:add(w6,[v,x['account'],issue['type'],issue['detail'],'AI_REVIEW',link(B/'analysis.json','逐条证据')])
    mapping={w7.cell(i,1).value:i for i in range(6,w7.max_row+1)}
    edits={'全项目状态':('PARTIAL','本条文字1通过，商品/史实待核；旧A83文字/25代表状态不变'),
        '更新于':(now(),'北京时间；用户指定单条，未重盘新增下载目录'),
        '来源视频':(416,'累计纳入415＋本条1；不是本次全目录盘点数'),'SHA唯一内容':(401,'本条SHA不在原400中；旧来源映射保留'),
        '本地转录/时间戳':(416,'本条新1404字/127段；旧415按历史完成记录'),
        '基础截图':(1664,'416来源×4基础帧；本条另10补帧，不混计'),
        '代表深度样本':(36,'旧35＋用户指定本条1；本条商品关键项AI_REVIEW'),
        'AI_REVIEW':(5,'已纳入36条深拆中的5待核；本批新增1，旧4保留'),
        '来源AI_PENDING':(297,'总览另有88个AI_REVIEW来源（旧87＋本条1）和31个限定范围AI_OK'),
        '全量文字语义完成':(1,'本条Codex全文127段及证据已核；旧400未验收'),
        '全量文字语义待做':(400,'旧317待做＋旧新两账号83待复核；总语料401'),
        '源视频不变':('本条1/1','本次仅核本条SHA/大小/mtime；旧415见10月16日历史校验'),
        '当前协作状态':('单条研究完成，商品待核','本条未派A；A旧5条修订不在本次审核范围'),
        '来源AI_REVIEW（总览）':(88,'旧87＋本条1；来源AI_OK仍31，AI_PENDING仍297')}
    for key,(value,note) in edits.items():i=mapping[key];put(w7,i,2,value);put(w7,i,3,note)
    add(w7,['本条本地处理耗时(秒)',read(B/'manifest/local_complete.json')['elapsed_s'],'离线CPU转录/4帧及冷启动；不是模型阅读耗时'])
    add(w7,['本条定向补帧',10,'加4基础帧，共14张实际查看；声音/连续运动未核'])
    for ws in [w1,w2,w3]:put(ws,3,1,'累计纳入416来源/401内容；36条深拆31可用/5待核；本条全文1通过，旧A83/25待复核状态保留。')
    put(w4,3,1,'新增用户指定单条；账号只计已纳入来源，全库频次保留历史400内容口径；未重盘其他下载。')
    for ws in [w1,w2,w3,w4,w6,w7]:
        ws.auto_filter.ref=f'A5:{get_column_letter(ws.max_column)}{ws.max_row}'
        for table in ws.tables.values():
            table.ref=ws.auto_filter.ref
            if table.autoFilter:table.autoFilter.ref=table.ref
    def check(book):
        n=0
        for title,cells in originals.items():
            for pos,(value,href,style) in cells.items():
                if (title,pos) in allowed:continue
                c=book[title][pos];assert c.value==value and (c.hyperlink.target if c.hyperlink else None)==href and c._style==style,(title,pos);n+=1
        return n
    preserved=check(wb);tmp=out.with_name(out.stem+'.single.tmp.xlsx');wb.save(tmp);saved=load_workbook(tmp)
    assert check(saved)==preserved and saved.sheetnames==wb.sheetnames
    counts=collections.Counter(saved.worksheets[0].cell(i,7).value for i in range(6,saved.worksheets[0].max_row+1));assert counts=={'AI_OK':31,'AI_REVIEW':88,'AI_PENDING':297}
    for ws in saved:
        for row in ws:
            for c in row:
                if (ws.title,c.coordinate) in allowed:assert c.data_type!='f'
                if c.hyperlink and c.hyperlink.target.startswith('file:'):assert pathlib.Path(unquote(urlparse(c.hyperlink.target).path).lstrip('/')).exists()
    assert sha(out)==sha(backup);tmp.replace(out)
    write(marker,{'at':now(),'status':'PASS','sha256':sha(out),'workbook_sha256':sha(out),'before_sha256':sha(backup),'prior_workbook_sha256':sha(backup),'unchanged_cells_verified':preserved,'source_rows':416,
                  'source_status_counts':dict(counts),'text_semantic_verified':1,'representative_scope':36,'deep_AI_OK':31,'deep_AI_REVIEW':5,'visual_qa':'PENDING'})
    print('SINGLE_WORKBOOK_UPDATED',preserved)

def record():
    done=B/'manifest/project_recorded.json'
    if done.exists():print('SINGLE_RECORD_RESUME_SKIP');return
    x=source();v=x['video_id'];obj=read(B/'analysis.json');book=read(B/'manifest/workbook_import.json');assert book['visual_qa']=='PASS'
    at=now();report=R/'reports/单条拆解_食力过日子_带孙辈长见识_历史名城下集.md';statepath=R/'manifest/collaboration_state.json';state=read(statepath)
    assert state['source_paths']==415 and state['unique_contents']==400 and sha(R/'阅见生活_竞品视频分析_终版.xlsx')==book['sha256']
    write(B/'manifest/collaboration_before.json',state)
    def append_one(path,row,expected):
        raw=path.read_bytes();old=[json.loads(s) for s in raw.splitlines() if s.strip()]
        assert len(old)==expected and v not in {z['video_id'] for z in old}
        path.write_bytes(raw+(b'' if raw.endswith(b'\n') else b'\n')+(json.dumps(row,ensure_ascii=False)+'\n').encode('utf-8'))
        assert path.read_bytes().startswith(raw);return sha(pathlib.Path(path))
    row={**x,'local_status':'LOCAL_OK','asr_status':'OK','ai_status':'AI_REVIEW','text_semantic_status':'AI_OK',
         'transcript_path':str(B/'transcripts'/f'{v}.txt'),'segments_path':str(B/'segments'/f'{v}.aligned.json'),
         'handoff_path':str(B/'manifest/READY.json'),'frames':read(B/'frames'/v/'index.json'),'ai_latest_record':str(B/'analysis.json'),
         'additional_frames':read(B/'manifest/supplementary_frames.json'),'latest_local_batch':KEY}
    manifest_sha=append_one(R/'manifest/current_manifest.jsonl',row,415)
    index={'video_id':v,'account':x['account'],'sha256':x['sha256'],'all_source_mappings':[{'video_id':v,'account':x['account'],'source_path':x['source_path']}],
           'duration_s':x['duration_s'],'source_title':x['source_title'],'transcript_path':row['transcript_path'],'segments_path':row['segments_path'],
           'text_characters':1404,'opening_asr_excerpt':(B/'transcripts'/f'{v}.txt').read_text(encoding='utf-8')[:120],
           'book_rule_candidates':[],'rule_cluster_ids':[],'mentioned_book_from_full_text':obj['fields']['recommended_book'],
           'legacy_deep_status':'AI_PENDING','latest_deep_status':'AI_REVIEW','latest_deep_record':str(B/'analysis.json'),'text_semantic_status':'AI_OK',
           'semantic_output_path':str(B/'analysis.json'),'facts_verified':False,'sales_performance_verified':False,'latest_local_batch':KEY,
           'scope':'Codex full-text plus 14 static frames; source/product issues still under review'}
    library_sha=append_one(R/'content_library/文案索引.jsonl',index,400)
    source_index=read(R/'manifest/current_source_index.json');source_index.update(at=at,sources=416,unique_contents=401,
        last_single_file_manifest=str((B/'manifest/source.json').relative_to(R)),counts_scope='Cumulative registered sources, not an October 19 re-inventory of all downloads')
    source_index['accounts'][x['account']]=1;write(R/'manifest/current_source_index.json',source_index)
    state.update(updated_at=at,latest_batch=KEY,source_paths=416,unique_contents=401,deep_sample_count=36,deep_status_counts={'AI_OK':31,'AI_REVIEW':5},
        source_AI_status_counts=book['source_status_counts'],workbook_sha256=book['sha256'],workbook_validation=book,
        current_inventory_scope='Cumulative registered415 plus one user-requested source; other new downloads not inventoried')
    state['text_semantic'].update(verified_complete=1,unaccepted_unique=400,independent_Codex_single_reviews=1,
                                  prepared_A_scope_note='Existing429 units/55 batches belong to the prior400; the new single-file review did not dispatch A')
    state['current_account_source_counts'][x['account']]=1
    state['unresolved_samples'].append({'video_id':v,'account':x['account'],'issues':[i['detail'] for i in obj['issues']]})
    state['latest_source_integrity']={'at':at,'status':'PASS_FOR_REQUESTED_FILE','sources_checked':1,'source_sha_size_mtime_unchanged':1,
                                     'database_opened':False,'scope':'Only this source; previous415 not rehashed in this review'}
    state['single_file_reviews']=[{'batch_id':KEY,'video_id':v,'source_account':x['account'],'text_verified':1,'deep_AI_REVIEW':1,'report':str(report)}]
    state['entry_paths']['latest_report']=str(report)
    # Keep the prior A review and its five-file repair task as historical, unaccepted work.
    state['next_Codex']='本条单文件研究已完成并登记；实际商品/史实待核。旧A的5条修订尚需另行验收，本轮未读取或接续'
    write(statepath,state)
    current=R/'reports/最终验收报告.md';archive=R/'reports/history/最终验收报告_20261019单条前.md'
    if not archive.exists():shutil.copy2(current,archive)
    oldissues='\n'.join(f'- `{d["video_id"]}`（{d["account"]}）：'+ '；'.join(d['issues']) for d in state['unresolved_samples'])
    current.write_text(f'''# 阅见生活竞品分析｜当前最终验收报告

更新于：{at}。**全项目PARTIAL；本轮为用户指定单条研究，未重盘整个下载目录、未接续旧A批次。**

## 当前累计口径

- 已纳入416来源、401个SHA独立内容（原415/400＋本条1/1）；这不是10月19日全目录实际文件数。
- 本地基础按416来源记录，新增一条1404字/127段、4基础帧及10补帧；旧415按原时点完成记录，不冒充本轮重验。
- 已纳入深拆36条：旧35＋本条1，31限定范围可用、5待核。总览来源31 AI_OK、88 AI_REVIEW、297 AI_PENDING。
- 独立全文语义验收1/401：本条已完整阅读并核证据，旧400中317待做、83待复核。全文理解不等同于史实、商品或营销效果验证。
- 旧A的83文字/25代表维持未通过；10月16日交回的5条修订本轮未读取或验收，不能视为当前仍在运行或已完成。
- 6套既有结构模板及其他原创产物沿用各自任务记录，本轮未新增原创成稿或发布。

## 本轮单条

[食力过日子：带孙辈长见识的历史名城下集]({report.name})。127段全文/14帧已查看，格式、引句、时间与身份通过；商品封面不符、内容承诺缺口及史实/活动声称仍待核。明确图书销售段112.235秒（34.3%），不能据此推断成交。

## 当前未决样本

{oldissues}

## 验证边界与入口

本次只重算目标源视频SHA/大小/mtime，一致；下载数据库未打开。旧415源视频与旧A证据的不变性只引用10月16日历史报告，前轮WAL/SHM变化原因仍未确认。表格保存前后核对{book['unchanged_cells_verified']}个非授权变更单元格的值/样式/链接；6个改动范围预览通过。渲染器保存全部图后退出码1，PNG与XLSX另行检查；未做原生Excel交互验收。

- [同一7-Sheet总表](../阅见生活_竞品视频分析_终版.xlsx)
- [401份文案索引](../content_library/文案索引.jsonl)
- [本条完整转录](../increments/{KEY}/transcripts/{v}.timed.txt)
- [旧A的07:38审核报告](new_accounts_2026-10-16_0738_A审核报告.md)
- [协作总控](../../docs/竞品分析协作总控.md)
''',encoding='utf-8',newline='\n')
    page=R/'content_library/开始使用.md';text=page.read_text(encoding='utf-8')
    text=re.sub(r'更新于：[^。]+。',f'更新于：{at}。',text,count=1)
    text=text.replace('400份独立内容全文已转录并建立索引（旧317＋新增83）；全量语义验收仍为0','401份已纳入独立内容全文可查（历史400＋用户指定1）；独立全文语义验收1/401')
    text=text.replace('已合并35份代表样本：31限定范围可用、4待核','已纳入36份代表样本：31限定范围可用、5待核')
    text+=f'\n## 2026-10-19 用户指定单条\n\n[带孙辈长见识：历史名城下集](../reports/{report.name})已完成全文与14帧研究。文字理解1通过，商品/史实AI_REVIEW；原A83/25未验收状态保留。\n'
    page.write_text(text,encoding='utf-8',newline='\n')
    control=ROOT/'docs/竞品分析协作总控.md';text=control.read_text(encoding='utf-8')
    text=re.sub(r'更新于：[^。]+。',f'更新于：{at}。',text,count=1)
    text=text.replace('|原素材|415来源／400SHA内容；旧332＋新增83来源，全量本地基础就绪|','|已纳入素材|416来源／401SHA内容；本轮只补用户指定1条，其他新下载未盘点|')
    text=text.replace('|代表深拆|35条（原25＋新10），31条限定范围可用、4条待核|','|代表深拆|36条（原35＋本条1），31条限定范围可用、5条待核|')
    text=text.replace('|全量文字轻分析|0/400已验收；旧317待做；新83（101单元/17批）已收但均AI_REVIEW，当前只修5条|','|全量文字轻分析|1/401已验收（本条）；旧317待做，83待复核。旧5条修订本轮未验收|')
    text=text.replace('## 新两账号｜最新审核状态','## 新两账号｜2026-10-16审核记录')
    text+=f'''\n## 2026-10-19｜用户指定单条：食力过日子

完成[历史名城下集单条研究](../竞品视频分析_2026-10-15/reports/{report.name})：127段全文与14静态帧。独立全文1通过，商品封面与史实/活动等AI_REVIEW；不认定播放/转化效果，不因本条覆盖旧A待办。

累计已纳入416来源/401独立；36条深拆31可用/5待核，全文语义1/401。仅新增指定文件，不宣称9账号全部下载数；全库频次/相似聚类保留旧400内容的历史计算，未重算。

本条未派A、无新自动任务或收费API；原视频/数据库只读。旧A的5条修订没有在本轮读取或确认状态，原工作范围与待验收记录保留。
'''
    control.write_text(text,encoding='utf-8',newline='\n')
    page=R/'README.md';text=page.read_text(encoding='utf-8');text=re.sub(r'更新于：[^。]+。',f'更新于：{at}。',text,count=1)
    text=text.replace('415来源／400SHA基础处理完成；35条代表样本整合，31条限定范围可用、4条待核','累计纳入416来源／401SHA；36条代表样本，31条限定范围可用、5条待核')
    text=text.replace('全文语义验收0/400：旧317待做、新83待复核','全文语义验收1/401（本条）：旧317待做、83待复核')
    text+=f'\n本轮单条：[带孙辈长见识](reports/{report.name})。只处理指定文件，其他新下载未盘点；当前总控与机器状态优先于历史段落。\n'
    page.write_text(text,encoding='utf-8',newline='\n')
    event={'batch_id':KEY,'at':at,'kind':'user_requested_single_file','source_added':1,'unique_added':1,'text_verified_added':1,'deep_review_added':1,
           'source_paths':416,'unique_contents':401,'deep_cumulative':36,'deep_status_counts':{'AI_OK':31,'AI_REVIEW':5},'project_status':'PARTIAL','report':str(report),'workbook_sha256':book['sha256']}
    log=R/'manifest/analysis_batches.jsonl';raw=log.read_bytes();assert not any(json.loads(s).get('batch_id')==KEY for s in raw.splitlines() if s.strip());log.write_bytes(raw+(json.dumps(event,ensure_ascii=False)+'\n').encode('utf-8'))
    journal=ROOT/'05_爆款知识库与复盘/06_竞品分析更新日志.md'
    with journal.open('a',encoding='utf-8',newline='\n') as f:f.write(f'\n## {at}｜用户指定历史名城下集单条\n\n食力过日子1条/327.634秒，全文127段与14帧研究完成；文字1通过，商品/史实AI_REVIEW。明确售书段112.235秒、34.3%，发现前五承诺与详细展开不符、示意封面不符口播商品。累计纳入416来源/401独立、36深拆31可用/5待核；原A83/25状态保留，未派A或重盘其他下载。[单条报告](../竞品视频分析_2026-10-15/reports/{report.name})。\n')
    write(done,{'at':at,'source_prefix_415_preserved':True,'library_prefix_400_preserved':True,'source_manifest_sha256':manifest_sha,'library_sha256':library_sha,
                'project_status':'PARTIAL','text_verified_added':1,'deep_status':'AI_REVIEW','report_sha256':sha(report),'workbook_sha256':book['sha256']})
    print('SINGLE_PROJECT_RECORDED')

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('phase',choices=['analysis','workbook','record']);args=ap.parse_args();{'analysis':analysis,'workbook':workbook,'record':record}[args.phase]()
