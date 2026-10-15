"""V4 explicitly requests openpyxl; author one basic seven-sheet workbook."""
from __future__ import annotations
import collections, datetime, json, math, re
from urllib.parse import urlparse, unquote
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from pipeline import OUTPUT_DIR as R, now, atomic_json

NAMES=['01_视频总览','02_带货脚本拆解','03_视觉模板','04_账号对标与高频规律','05_原创带货模板库','06_异常待复核','07_口径与运行摘要']
ISSUES={'HOOK_0_3S_CONTAINS_AUDIO_AFTER_3S':'0～3秒口播字段包含3秒之后的内容','NATIVE_VIDEO_CAPABILITY_IS_A_SELF_REPORT_NOT_INDEPENDENTLY_VERIFIED':'原生视频理解能力仅有A自述，尚缺独立核验','UNSUPPORTED_PERFORMANCE_OR_ABSOLUTE_CLAIM':'含高完播/高口碑/100%等缺少依据的判断','EARLY_TRANSCRIPT_END_REVIEW':'转录结束早于视频结尾30秒以上，尾段待听核'}
def read(p,default=None):return json.loads(p.read_text(encoding='utf-8')) if p.exists() else default
def safe(v):
    if v is None:return 'AI_PENDING'
    if isinstance(v,str):
        v=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',v)
        return "'"+v if v.lstrip().startswith(('=','+','-','@')) else v
    return v
def link(path,label):return {'path':str(path),'label':label}
def compact(values,n=3):
    values=list(dict.fromkeys(values));return '；'.join(values[:n])+(f' 等{len(values)}项' if len(values)>n else '')
def table(wb,name,headers,rows,widths,subtitle):
    ws=wb.create_sheet(name);ws.sheet_view.showGridLines=False;ws.sheet_view.zoomScale=85;ws.sheet_properties.tabColor='A9CDE8'
    ws.cell(2,1,name[3:]);ws.cell(2,1).font=Font(name='Microsoft YaHei',size=16,bold=True,color='234B68')
    ws.cell(3,1,subtitle);ws.cell(3,1).font=Font(name='Microsoft YaHei',size=10,color='5C6570')
    for j,h in enumerate(headers,1):
        c=ws.cell(5,j,h);c.font=Font(name='Microsoft YaHei',bold=True,size=10,color='234B68');c.fill=PatternFill('solid',fgColor='DCECF7');c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width=widths[j-1] if j<=len(widths) else 24
    ws.row_dimensions[5].height=32
    for i,row in enumerate(rows,6):
        for j,val in enumerate(row,1):
            if isinstance(val,dict) and 'path'in val:
                c=ws.cell(i,j,val['label']);p=Path(val['path']);c.hyperlink=p.as_uri();c.font=Font(name='Microsoft YaHei',size=10,color='0563C1',underline='single')
            else:
                c=ws.cell(i,j,safe(val));c.font=Font(name='Microsoft YaHei',size=10,color='273746')
                if isinstance(val,str):c.number_format='@'
                if isinstance(val,(int,float)):c.number_format='#,##0.0' if isinstance(val,float) else '#,##0'
            c.alignment=Alignment(vertical='top',wrap_text=True,horizontal='right' if isinstance(val,(int,float)) else 'left')
            if isinstance(val,str) and val in ['AI_PENDING','AI_REVIEW','LOCAL_FAILED','NO_SPEECH']:
                c.fill=PatternFill('solid',fgColor='FFF2CC' if val in ['AI_PENDING','AI_REVIEW'] else 'FCE4D6')
        ws.row_dimensions[i].height=48 if name not in ['05_原创带货模板库','07_口径与运行摘要'] else 42
    last=max(6,5+len(rows));ref=f'A5:{get_column_letter(len(headers))}{last}';ws.auto_filter.ref=ref;ws.freeze_panes='C6' if len(headers)>5 else 'A6'
    if rows:
        tbl=Table(displayName='Data'+name[:2],ref=ref);tbl.tableStyleInfo=TableStyleInfo(name='TableStyleLight9',showRowStripes=True);ws.add_table(tbl)
    ws.print_options.horizontalCentered=False;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:5'
    return ws

def main():
    rows=[json.loads(s) for s in (R/'manifest/manifest.jsonl').read_text(encoding='utf-8').splitlines() if s];lookup={x['video_id']:x for x in rows}
    plan=read(R/'manifest/phase2_plan.json');summary=read(R/'manifest/phase2_analysis_summary.json');integrity=read(R/'manifest/phase2_source_integrity.json',{})
    features=read(R/'clusters/video_features.json',{});exact=read(R/'clusters/exact_duplicates.json',[]);textclusters=read(R/'clusters/text_clusters.json',[]);books=read(R/'clusters/book_frequencies.json',[]);people=read(R/'clusters/story_entity_frequencies.json',[]);candidates=read(R/'clusters/deep_candidates.json',[])
    av=read(R/'manifest/ai_validation.json',{'records':[]})['records'];ai={z['video_id']:z for z in av if z['schema_valid']};members=collections.defaultdict(list)
    for c in exact+textclusters:
        for v in c['video_ids']:members[v].append(c['cluster_id'])
    wb=Workbook();wb.remove(wb.active);overview=[];scripts=[];visual=[];exceptions=[];canonstates={};ready=0
    for x in rows:
        v=x['video_id'];canonical=plan['canonical_by_video_id'][v];feat=features.get(canonical,{});st=read(R/'state'/f'{v}.json',{});frames=read(R/'frames'/v/'index.json',[]);aligned=read(R/'segments'/f'{v}.aligned.json',{'segments':[]});tx=R/'transcripts'/f'{v}.txt';airow=ai.get(v);aistat=airow['import_status'] if airow else 'AI_PENDING';fields=airow['fields'] if airow and aistat=='AI_OK' else {}
        astep=st.get('steps',{}).get('asr',{});asr=astep.get('status','PENDING');canonstates[canonical]=read(R/'state'/f'{canonical}.json',{})
        if(R/'handoff'/v/'READY.json').exists():ready+=1
        title=((x.get('db')or{}).get('title') or Path(x['source_path']).stem).split('#')[0].strip()
        overview.append([x['account'],v,title[:90],x.get('duration_s',0),st.get('status','DISCOVERED'),asr,aistat,canonical,compact(feat.get('book_candidates',[])) or '未命中规则词典',bool(feat.get('book_mention_rule')),bool(feat.get('recommendation_rule')),bool(feat.get('shopping_prompt_rule')),compact(members[v]+([] if canonical==v else members[canonical])) or '无阈值簇',next((c['rank']for c in candidates if c['video_id']==canonical),'未选'),x['size_bytes'],x['sha256'],str((x.get('db')or{}).get('video_id','未关联')),link(x['source_path'],'原视频'),link(tx,'全文TXT'),link(R/'segments'/f'{v}.aligned.json','时间戳JSON'),link(R/'handoff'/v/'input.json','交接包'),fields.get('theme','AI_PENDING'),next((link(R/'clusters/deep_candidates.json',c['reason'][:80]+'（完整见链接）')for c in candidates if c['video_id']==canonical),'未入选（保留基础处理）')])
        hook='0～3秒词级对齐待复核，见分段JSON'
        # Exact local 0-3 seconds uses tokens, not whole first sentence.
        raw=read(R/'segments'/f'{v}.json',{});t=tx.read_text(encoding='utf-8').strip() if tx.exists() else ''
        try:
            from align_segments import token_spans
            ts=token_spans(t);tim=raw.get('raw_timestamp',[])
            if len(ts)!=len(tim):ts=token_spans(t,keep_apostrophe=True)
            if len(ts)==len(tim):hook=''.join(a[2] for a,b in zip(ts,tim) if b[0]<3000) or '0～3秒无ASR词元（待核原音）'
        except Exception:pass
        scripts.append([x['account'],v,aistat,hook[:120],fields.get('hook',{}).get('hook_type','AI_PENDING'),fields.get('theme','AI_PENDING'),fields.get('narrative_structure','AI_PENDING'),fields.get('recommended_book','AI_PENDING'),fields.get('bridge',{}).get('text','AI_PENDING'),fields.get('bridge',{}).get('start_time_s','AI_PENDING'),fields.get('cta',{}).get('text','AI_PENDING'),fields.get('cta',{}).get('start_time_s','AI_PENDING'),feat.get('book_marker_excerpt','未命中规则词')[:90],feat.get('book_marker_count',0),feat.get('purchase_marker_count',0),fields.get('target_audience','AI_PENDING'),link(R/'segments'/f'{v}.aligned.json','原话与时间'),link(airow['path'],'A原JSON') if airow else 'AI_PENDING',fields.get('audience_pain','AI_PENDING'),fields.get('value_proposition','AI_PENDING'),airow.get('input_mode','AI_PENDING') if airow and aistat=='AI_OK' else 'AI_PENDING'])
        rr=[x['account'],v]
        for i in range(4):
            f=frames[i]if i<len(frames)else{};rr+=[f.get('time_s','PENDING'),link(f['path'],f'帧{i+1}')if f else 'PENDING']
        va=fields.get('visual_and_audio',{});rr+=[va.get('visual_type','AI_PENDING'),va.get('subtitle_style','AI_PENDING'),va.get('book_appearance','AI_PENDING'),va.get('audio_style','AI_PENDING'),aistat,'固定4帧不能证明完整镜头和音乐节奏'];visual.append(rr)
        if asr not in ['OK']or st.get('status')=='LOCAL_FAILED':exceptions.append([v,x['account'],'本地处理',st.get('error',asr),'待处理',link(R/'state'/f'{v}.json','状态证据')])
        if aligned.get('alignment_method')=='CHUNK_LEVEL_FALLBACK':exceptions.append([v,x['account'],'时间戳降级',aligned.get('alignment_warning',''),'仅保留真实分段区间；未造词级时间',link(R/'segments'/f'{v}.aligned.json','时间戳证据')])
    table(wb,NAMES[0],['账号','video_id','来源标题','时长(秒)','本地状态','ASR状态','AI状态','相同SHA复用ID','候选书名（非售书结论）','书名提及规则','明确推荐词规则','购物引导词规则','重复/相似簇','深度候选序号','大小(字节)','SHA-256','平台视频ID','原视频','转录全文','分段时间戳','READY交接包','AI主题','深度候选原因'],overview,[22,25,56,14,22,14,16,25,36,18,18,18,36,15,17,26,24,12,12,15,15,36],'Phase 2 基础版｜规则命中只作候选；AI_PENDING/AI_REVIEW均不计作AI已完成。')
    table(wb,NAMES[1],['账号','video_id','AI状态','本地ASR 0～3秒（词起点<3秒）','Hook类别','主题','叙事结构','实际推荐书名','Bridge原话','Bridge起点(秒)','CTA原话','CTA起点(秒)','本地荐书词命中片段（待复核）','荐书词次数','购买词次数','受众','原话时间证据','A原JSON（待校验状态见摘要）','受众痛点','产品价值锚点','实际输入模式'],scripts,[22,25,16,52,25,30,45,24,46,16,46,16,55,14,14,34,16,36],'本地ASR摘录不是人工核实引语；AI未通过验收的业务字段保持AI_PENDING。')
    table(wb,NAMES[2],['账号','video_id','帧1秒','帧1','帧2秒','帧2','帧3秒','帧3','帧4秒','帧4','画面形式','字幕样式','书籍露出','配音/音乐','AI状态','局限'],visual,[22,25,13,12,13,12,13,12,13,12,35,32,36,35,16,44],'4帧固定位置：1秒、5秒、50%、倒数2秒；时间均为实际使用值。')
    freq=[]
    for a,n in collections.Counter(x['account']for x in rows).items():
        unique=len({x['sha256']for x in rows if x['account']==a});nc=sum(c['account']==a for c in candidates);freq.append(['账号',a,n,unique,1,nc,'同账号SHA去重；非语义唯一数',link(R/'manifest/manifest.jsonl','来源清单')])
    for b in books:freq.append(['书名候选',b['term'],b['video_count'],b['video_count'],b['account_count'],compact(e['video_id']for e in b['evidence']),'提及≠推荐；正式书名/商品待核',link(R/'clusters/book_frequencies.json','逐条证据')])
    for p in people:freq.append(['人物/故事',p['term'],p['title_focus_count'],p['video_count'],p['account_count'],compact(e['video_id']for e in p['evidence']),'列C=标题主讲数；列D=全文或标题提及数',link(R/'clusters/story_entity_frequencies.json','逐条证据')])
    for k in read(R/'clusters/keyword_frequencies.json',[])[:30]:freq.append(['文本关键词候选',k['keyword'],k['video_count'],k['video_count'],k['account_count'],compact(e['video_id']for e in k['evidence']),'名词TF-IDF候选；分词/主题未人工复核',link(R/'clusters/keyword_frequencies.json','逐条证据')])
    for c in exact+textclusters:freq.append([c['kind'],c['cluster_id'],len(c['video_ids']),len(c['video_ids']),c['account_count'],compact(c['video_ids']),'连通分量≠任意两条皆相似；配对分数另见JSON',link(R/'clusters'/('exact_duplicates.json' if c['kind']=='SHA256_EXACT' else 'text_clusters.json'),'簇证据')])
    table(wb,NAMES[3],['类型','账号/词项/簇','来源数或标题主讲数','去重内容/提及数','账号数','候选数或证据video_id','口径','证据'],freq,[22,32,20,20,14,58,58,16],'不同SHA内容为频次分母；频次仅用于研究候选，不证明高播放或销量。')
    table(wb,NAMES[4],['模板ID','状态','适用受众','开头结构','故事结构','Bridge','CTA','视觉形式','参考ID/候选清单','说明'],[['AI_PENDING','AI_PENDING','AI_PENDING','AI_PENDING','AI_PENDING','AI_PENDING','AI_PENDING','AI_PENDING',link(R/'clusters/deep_candidates.json','25条深度候选'),'Phase 3完成且证据核验后归纳3～8套；基础版不凑模板。']],[24,18,30,35,35,35,35,32,24,68],'原创模板尚未完成；单条A建议不能直接当作跨样本模板结论。')
    for z in av:
        for issue in z['schema_errors']+z['content_issues']:exceptions.append([z.get('video_id','未知'),lookup.get(z.get('video_id'),{}).get('account','未知'),'A结果校验',ISSUES.get(issue,issue),z['import_status'],link(z['path'],'A原JSON')])
    for e in read(R/'reports/专有名词候选词典.json',{}).get('entries',[]):exceptions.append([e['video_id'],lookup[e['video_id']]['account'],'专名候选修订',e['asr']+' → '+e['candidate'],e['review'],link(R/'reports/专有名词候选词典.json','候选词典')])
    for q in read(R/'manifest/asr_quality_flags.json',{'records':[]})['records']:exceptions.append([q['video_id'],q['account'],'ASR质量规则触发','；'.join(ISSUES.get(f,f)for f in q['flags']),'待人工/A点检；不等于素材损坏',link(q['evidence_path'],'原话时间证据')])
    for fp in sorted((R/'logs/failures').glob('*.json')):
        h=read(fp,{});v=h.get('video_id','未知');current=read(R/'state'/f'{v}.json',{})
        description='Windows状态文件短暂被占用（拒绝访问），原始错误见链接'if 'WinError 5'in h.get('error','')else h.get('error','见日志')[:110]
        exceptions.append([v,lookup.get(v,{}).get('account','未知'),'已记录的历史失败',description,'补跑后成功' if current.get('steps',{}).get('asr',{}).get('status')=='OK' else '尚未恢复',link(fp,'原始失败记录')])
    exceptions.append(['全部','全部','ASR质量边界','未经逐句耳听；有些词名和历史叙述可能有误','等待A/人工抽查',link(R/'reports/10条样本验收.md','样本质量记录')])
    table(wb,NAMES[5],['video_id','账号','问题类型','问题','状态/处理方式','证据'],exceptions,[25,22,26,74,44,22],'保留异常与未完成项；未用健康零值掩盖待复核。')
    job=read(R/'state/phase2_job.json',{});ast=sum(s.get('steps',{}).get('asr',{}).get('elapsed_s',0) or 0 for s in canonstates.values())
    metrics=[['阶段','Phase 2 基础版','AI深度/原创模板仍待完成'],['导出时间',now(),'本地执行时区+08:00'],['来源路径',len(rows),'一文件一行'],['SHA唯一内容',len(set(plan['canonical_by_video_id'].values())),'不同编码相似内容未合并为同一文件'],['完全重复额外文件',summary['extra_exact_duplicates'],'保留所有来源'],['有转录路径',sum((R/'transcripts'/f'{x["video_id"]}.txt').exists()for x in rows),'包含SHA复用'],['READY交接包',ready,'最后写READY.json'],['图片数',sum(len(list((R/'frames'/x['video_id']).glob('*.jpg')))for x in rows),'应为4×来源文件数'],['AI Schema通过',summary['ai_schema_valid'],'不等于内容通过'],['AI_OK',summary['ai_ok'],'只计证据与内容均通过'],['AI_REVIEW',summary['ai_review'],'核心AI字段保留AI_PENDING'],['AI_PENDING',len(rows)-summary['ai_ok']-summary['ai_review'],'尚无通过结果'],['全文相似簇',summary['text_clusters'],'全文TF-IDF≥0.76、长度比≥0.6、顺序匹配≥0.65'],['开头相似簇',summary['opening_clusters'],'前100字字符TF-IDF≥0.8'],['段落复用簇',summary['passage_clusters'],'公共片段可能是故事或通用CTA'],['相似配对',summary['similarity_pairs'],'保留每对实际证据'],['跨账号相似配对',summary['cross_account_pairs'],'不是销量指标'],['深度候选',len(candidates),'按SHA唯一内容计数；重复来源共享序号'],['最近一次补跑墙钟(秒)',job.get('elapsed_s','运行中'),'仅最终单进程补跑；全轮累计另列'],['规范ID状态所记ASR累计(秒)',ast,'含样本缓存；SHA复用可能记0，不是墙钟时间'],['源文件复核不变数',integrity.get('unchanged_files','待复核'),'按全部SHA前后比对'],['数据库','本轮不打开','沿用Phase0只读取得字段；前轮物理变化已记录'],['FunASR','1.4.16','本地缓存模型，无收费API'],['PyTorch','2.14.1+cpu','CPU运行'],['Python','3.10.10','分析/ASR独立环境'],['FFmpeg','9.0.2','本机输出目录工具'],['外部文本','防公式注入','危险前导字符写成文本；ID使用文本'],['A文件导入规则','稳定完整JSON+Schema+证据校验','不改analysis_a；缺失/待复核AI字段=AI_PENDING']]
    metrics.append(['ASR质量复核提示',read(R/'manifest/asr_quality_flags.json',{}).get('flagged_paths','待检查'),'规则触发，不等于损坏；转录生成成功不等于逐字正确率'])
    validation=read(R/'manifest/phase2_validation.json',{})
    metrics.append(['累计处理墙钟(秒)',validation.get('run_wall_seconds_including_restart_test','待检查'),'含真实中断测试、重启与失败补跑；不含后续聚类/导出'])
    metrics.append(['新增派生素材(GiB)',validation.get('derived_files_bytes',0)/1024**3,'音频/文字/时间戳/截图/交接副本/分段；不含环境及模型'])
    metrics.append(['收费API调用',0,'本地处理；电费及设备折旧未计量'])
    table(wb,NAMES[6],['指标','实测/状态','定义与局限'],metrics,[35,46,85],'来源：manifest、state、clusters及完整A文件的校验副本。')
    wb[NAMES[0]].column_dimensions['W'].width=68
    for row in wb[NAMES[0]].iter_rows(min_row=6,min_col=10,max_col=12):
        for cell in row:cell.value='命中'if cell.value else'未命中'
    out=R/'阅见生活_竞品视频分析_终版.xlsx';tmp=out.with_suffix('.tmp.xlsx');wb.save(tmp);tmp.replace(out)
    check=load_workbook(out,read_only=False,data_only=False);assert check.sheetnames==NAMES;assert check[NAMES[0]].max_row==len(rows)+5
    assert {check[NAMES[0]].cell(i,2).value for i in range(6,len(rows)+6)}==set(lookup)
    for i,x in enumerate(rows,6):
        expected=str((x.get('db')or{}).get('video_id','未关联'))
        cell=check[NAMES[0]].cell(i,17)
        assert cell.value==expected and cell.data_type=='s' and cell.number_format=='@',f'Platform ID corrupted: {x["video_id"]}'
    errors=[];links=0
    for ws in check:
        assert ws.freeze_panes and ws.auto_filter.ref
        for rr in ws.iter_rows():
            for c in rr:
                if c.data_type=='f':errors.append(f'Unexpected formula {ws.title}!{c.coordinate}')
                if c.hyperlink:
                    links+=1;u=urlparse(c.hyperlink.target)
                    if u.scheme=='file':
                        target=unquote(u.path)
                        if re.match(r'^/[A-Za-z]:/',target):target=target[1:]
                        if not Path(target).exists():errors.append(f'Missing hyperlink target {ws.title}!{c.coordinate}: {target}')
    assert not errors,errors
    atomic_json(R/'manifest/workbook_validation.json',{'at':now(),'path':str(out),'sheets':NAMES,'video_rows':len(rows),'unique_video_ids':len(lookup),'hyperlinks':links,'formula_injection_scan':'PASS','freeze_filter':'PASS','schema_counts':summary,'status':'STRUCTURAL_PASS_VISUAL_QA_PENDING'})
    print(str(out),flush=True)

if __name__=='__main__':main()






