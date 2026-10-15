import collections, json, datetime, re
from pathlib import Path
from PIL import Image
from openpyxl import load_workbook
from pipeline import OUTPUT_DIR as R, atomic_json, atomic_text, now, sha256
from phase2_export import NAMES

def read(p):return json.loads(p.read_text(encoding='utf-8'))
rows=[json.loads(s)for s in(R/'manifest/manifest.jsonl').read_text(encoding='utf-8').splitlines()if s];lookup={x['video_id']:x for x in rows};records={p.stem:read(p)for p in(R/'manifest/phase4_records').glob('v_*.json')};summary=read(R/'manifest/phase4_import_summary.json');integrity=read(R/'manifest/phase4_source_integrity.json');export=read(R/'manifest/phase4_export_state.json');accounts=read(R/'manifest/phase4_account_summary.json')['accounts'];templates=read(R/'manifest/phase4_templates.json')['templates']
out=R/'阅见生活_竞品视频分析_终版.xlsx';wb=load_workbook(out,data_only=False)
assert sha256(out)==export['output_sha256']and wb.sheetnames==NAMES and len(list(R.glob('*.xlsx')))==1
expected={c['video_id']for c in read(R/'clusters/deep_candidates.json')};assert set(records)==expected and len(records)==25
assert summary['schema_pass']==25 and integrity['status']=='PASS'
assert read(R/'manifest/phase4_resume_result.json')['status']=='PASS'and read(R/'manifest/phase4_interruption_exit.json')['actual_child_exit_code']==75
assert read(R/'manifest/phase4_export_resume_result.json')['hash_mtime_unchanged']
frames=0
for i,x in enumerate(rows,6):
    v=x['video_id'];src=Path(x['source_path']);st=src.stat();assert st.st_size==x['size_bytes']and st.st_mtime_ns==x['mtime_ns']
    for ws in wb.worksheets[:3]:assert ws.cell(i,2).value==v
    assert wb.worksheets[0].cell(i,17).value==str((x.get('db')or{}).get('video_id','未关联'))
    t=(R/'transcripts'/f'{v}.txt').read_text(encoding='utf-8').strip();segments=read(R/'segments'/f'{v}.aligned.json');assert ''.join(s['text']for s in segments['segments'])==t
    ff=read(R/'frames'/v/'index.json');assert len(ff)==4
    for f in ff:
        with Image.open(f['path'])as im:im.verify()
        frames+=1
    assert(R/'handoff'/v/'READY.json').exists()
    if v in records:
        r=records[v];n=r['normalized'];assert wb.worksheets[1].cell(i,3).value==r['status'];assert wb.worksheets[1].cell(i,4).value==n['hook_text'];assert wb.worksheets[1].cell(i,8).value==n['recommended_book']
        for k in ['bridge','cta']:
            z=n[k]
            if z['present']:assert z['text']in t and 0<=z['start_s']<z['end_s']<=x['duration_s']+1
    else:assert wb.worksheets[0].cell(i,7).value=='AI_PENDING'
assert frames==1328
for z in read(R/'manifest/phase4_input_snapshot.json')['files']:assert sha256(Path(z['source']))==z['sha256']and Path(z['source']).stat().st_mtime_ns==z['mtime_ns']
states=collections.Counter(wb.worksheets[0].cell(i,7).value for i in range(6,338));assert states=={'AI_OK':20,'AI_REVIEW':5,'AI_PENDING':307}
assert wb.worksheets[4].max_row==10
for i,a in enumerate(accounts,6):assert [wb.worksheets[3].cell(i,j).value for j in range(5,10)]==[a['sample_count'],a['ai_ok'],a['ai_review'],a['book_introduction_count'],a['shopping_guidance_count']]
visual=read(R/'manifest/phase4_visual_qa.json')
for z in visual['previews']:
    with Image.open(z['path'])as im:im.verify()
assert len(visual['previews'])==10
render_exit=read(R/'manifest/phase4_render_process.json')['actual_process_exit_code']
visual.update({'reviewed_at':now(),'reviewer':'Codex','status':'PASS','sheets_visually_reviewed':7,'detail_previews_reviewed':3,'scope':'Seven sheet layouts and imported deep-sample cells inspected; final changed pending/data fills checked. Does not claim every cell was manually read.','render_process_exit_code':render_exit,'tool_limitation':'Renderer wrote all ten images, then Windows process exited abnormally. Images independently decoded with Pillow and actually reviewed. Numeric-string platform IDs may display scientifically in this renderer; XLSX text values independently matched exactly.'})
atomic_json(R/'manifest/phase4_visual_qa.json',visual)
export.update({'status':'PASS','visual_qa_reference':str(R/'manifest/phase4_visual_qa.json'),'scope':'Workbook structure, exact IDs, links, statuses, evidence and visual layout only; not full project acceptance'})
atomic_json(R/'manifest/phase4_export_state.json',export)
old=R/'manifest/workbook_validation.json';oldbackup=R/'manifest/phase2_workbook_validation_snapshot.json'
if not oldbackup.exists():oldbackup.write_bytes(old.read_bytes())
atomic_json(old,export)
pending=[{'video_id':v,'account':r['account'],'issues':r['unresolved_issues']}for v,r in records.items()if r['status']=='AI_REVIEW']
result={'at':now(),'project_status':'PARTIAL','phase4_integration_status':'PASS','files':332,'sha_unique_contents':317,'local_success_paths':332,'basic_frames':1328,'ready_packages':332,'deep_json_loaded':25,'schema_pass':25,'ai_ok':20,'ai_review':5,'ai_pending_outside_deep_scope':307,'templates_loaded':5,'research_usable_templates':4,'template_drafts_pending_evidence':1,'pending_records':pending,'source_unchanged':332,'A_files_unchanged':28,'import_resume_test':'PASS','export_resume_test':'PASS','previous_ASR_chunks_unchanged':'20/20','hyperlinks_checked':export['hyperlinks'],'workbook_sha256':sha256(out),'render_process_exit_code':render_exit,'scope_note':'AI_OK is limited to checked imported text/model timestamps/static frames, with permitted unknown fields. ASR word accuracy, factual truth, full audio and sales performance are not claimed verified.'}
atomic_json(R/'manifest/phase4_final_acceptance.json',result)
paths={'工作簿':out,'最终验收报告':R/'reports/最终验收报告.md','导入更正与待补证清单':R/'reports/Phase4_导入更正与待补证清单.md','整理版模板库':R/'reports/原创带货模板库_验收整理版.md','机器验收记录':R/'manifest/phase4_final_acceptance.json'}
lines=['# 阅见生活竞品分析｜最终验收报告（V4.0）','',f'验收时间：{result["at"]}。','', '**全项目终验：PARTIAL。Phase 4 的数据整合与工作簿交付已完成；存在明确列出的证据和质量缺项，不能写全项目PASS。**','', '## 1. 完成范围与口径','', '- 源视频332个、SHA唯一内容317份；317份识别及同SHA复用覆盖332来源，保留全部video_id与来源。基础截图1328张、READY交接包332个。','- A的25份单视频JSON与25条候选完全一致；25份结构校验通过。账号总结6账号、原始模板5套均已读取并保留原件。','- 导入后20条AI_OK、5条AI_REVIEW；其他307条未安排深度分析，按V4范围保留AI_PENDING。307条不构成本次必须追加的全量深度任务。','- AI_OK的范围是核验后的文本、模型估计时间和所见四帧；不表示史实、商品真伪、全部音画或营销效果均被证明。未知正式书名、声音判断明确标注，不硬补。','- 5套模板全部整理进Sheet05，4套研究结构可用，1套单样本草案待补第二来源。均不等于可直接发布的文案。','', '## 2. 六账号对账','','|账号|来源文件|SHA唯一|深度样本|AI_OK|AI_REVIEW|引书/阅读|明确商品引导|','|---|---:|---:|---:|---:|---:|---:|---:|']
for a in accounts:
    source=[x for x in rows if x['account']==a['account']];lines.append(f'|{a["account"]}|{len(source)}|{len({x["sha256"]for x in source})}|{a["sample_count"]}|{a["ai_ok"]}|{a["ai_review"]}|{a["book_introduction_count"]}|{a["shopping_guidance_count"]}|')
lines+=['','引书/阅读与商品引导仅在25条深度样本中计数；借阅、关注或文化提及不等同购物。无播放、商品点击与成交数据，不能把比例当真实商业转化率。','', '## 3. 导入与更正','', '- A原件只读，冻结快照后再校验；不读取临时文件，25条身份、账号、源路径、Schema与分段证据逐条对齐。','- 0～3秒按本地词级时间戳重建，A原摘录及完整Hook另列。所有ASR引语标明为模型转录，未冒充逐句耳听。','- 林睿v_5ddb5a8ad7ccdb86在196.540秒开始介绍日历书，244.320秒起有“放下面点开/弄一本”线索；补回A漏掉的书籍、Bridge与CTA。257.498秒既有帧可见365日走遍中国字样。','- 锦绣v_ac516db0b5dced62在351.110秒开始引书，642.950秒有“大美中国左下方商品”原文；补回遗漏。','- v_7c9d47c6f0bcb117的515.690秒转录称“把话说的恰到好处”，未采用A的“闲谈高手类”作为正式书名；借阅对话不计购物促单。v_d3f3bf3980928b72在457.940秒介绍“中国历史名城”系列，正式品名/版本仍待核，不保留“洛阳历史大书”作确定名称。','- 文字+四帧输入不能支持完整配乐判断；核验栏写无法确认，A原画面/声音描述另列。人物或历史影像的真实性不由画面风格替代核验。','- Sheet04保留原A账号判断供对照，统计改按可定位的引书与商品引导复算。Sheet05采用抽象结构、新提议题及风险，不照搬原模板整段促单话术。','', '## 4. 为什么不是PASS','']
for p in pending:lines.append('- '+p['video_id']+'（'+p['account']+'）：'+'；'.join(p['issues']))
lines+=['- T03另引的“第1期山西古建”没有video_id，未计为有效第二来源。','- v_2fe34883cba71ae1末尾约30秒ASR未出文字；有低音量非零音频，尚待听核。既有专名候选词典也保留待核项。','- 原A能力验证报告的部分指定时刻细节过强：120秒实帧为题字特写，不能声称该帧看见老人伏案写字；林睿报告中的茶水等细节未获实帧支持。8张独立补帧可回看，仍不足以证明具体原生输入工具链和完整听音能力。','', '## 5. 验收门槛与证据','','|验收项|结果|证据/局限|','|---|---|---|','|来源完整盘点与映射|PASS|332来源/317SHA，目录成员一致|','|源文件不变性|PASS|332/332 SHA、大小和mtime一致；本轮未打开下载数据库，数据库文件元数据与Phase2相同|','|本地文字/时间戳/截图完整性|PASS|332条分段可重建全文、1328张JPEG逐文件解码通过|','|ASR可用性与专名质量|PARTIAL|10样本已由用户确认；未量化全量逐字准确率，尾段与候选词保留待核|','|重复/相似/选题分口径|PASS（文本检查）|15完全重复簇，全文9、开头5、公共段落10；124对相似，跨账号0；21对Codex文本复核记录可查，不冒充用户人工耳听|','|A模态与关键证据|PARTIAL|20条限定范围可用；3条native声明/细节待补、2条事实叙事未核|','|深度样本范围|PASS|25条覆盖6账号，每账号4～5条；没有把332条都当AI分析完成|','|模板证据与原创结构|PARTIAL|4套研究结构可用，1套草案缺第二来源；价格/功效/收益保证与照搬话术已剔除|','|7-Sheet Excel|PASS（交付）|单文件恰好7表；总览332行，25条分析字段已填；3346个链接可达，平台长ID逐项精确文本核对；七表及深度样本区域已查看|','|断点恢复|PASS|导入子进程真实退出75后恢复；跳过10个已完成记录，断点前5个检查点哈希/mtime 5/5不变；既有ASR20个分段仍未重写|','|重复导出|PASS|输入未变时跳过导出，工作簿SHA和mtime不变|','|A原件保护|PASS|28个原交付文件SHA与mtime保持不变，原件与导入更正分开|','|权限/费用|PASS（本轮）|本轮无新上传、无收费API调用、未装新框架；A既有费用/额度未知，不宣称零消耗|','', '## 6. 工具与预览局限','',f'预览引擎完成10张图后，Windows进程异常退出，实际退出码{render_exit}；图片经Pillow独立解码并逐表/重点单元格检查。工作簿已由openpyxl独立读取、计数、长ID及链接校验。预览引擎可能把长数字文本显示为科学计数；Excel原件仍为精确文本，不能从图片抄录平台ID。异常码和日志均保留，未将渲染进程错误冒充正常退出。','', '## 7. 后续原创生产建议','', '1. 优先用T01“沟通场景→表达练习”和T05“风景问题→阅读延伸”，各做2条原创小样；原始命中频次只能作为试验理由，不能作为销量证明。','2. 先确定真实在售版本、目录及两三页可展示内容，再围绕真实用途写独立脚本。可用新题“临时加任务怎样问清优先级”“沿一条河读三座城”；这些是新提议题，效果尚未测试。','3. 人物逝世叙事不进入事实型成稿；传记模板先核史料。价格锚点、医疗长寿与收益保证不直接沿用。','4. 使用自摄/已许可图书与风景素材、本人或许可音色；不复用竞品完整开头、对话、画面与声音。不把降相似度作为原创标准。','5. 首轮只改一个变量，例如同一书的两种开头；记录留存、完播、商品点击、成交及退款等实际指标后，再决定扩大。此处为生产试验建议，不自动发布。','', '## 8. 产物绝对路径','']
lines += [f'- {label}：[{p.name}]({p.as_posix()})'for label,p in paths.items()]
lines+=['',f'- A原件目录：{(R/"analysis_a").as_posix()}','- 输入快照、逐视频校验记录、账号复算与模板证据分别位于manifest/phase4_input_snapshot、phase4_records、phase4_account_summary.json、phase4_templates.json。','- 当前最终状态以本报告与manifest/phase4_final_acceptance.json为准；Phase2报告及其SHA是历史阶段记录。']
atomic_text(paths['最终验收报告'],'\n'.join(lines)+'\n')
tl=['# 原创带货模板库｜验收整理版','', '来源于A的5套结构，保留原件；本文件抽象方法与新提议题由Codex整理。4套研究结构可用、1套单样本草案。所有题目和时长均待测试，不等于发布成稿。','']
for t in templates:
    tl += [f'## {t["id"]} {t["name"]}｜{t["status"]}','',f'- 受众：{t["audience"]}；图书范围：{t["books"]}；时长：{t["seconds"]}',f'- 开头：{t["hook"]}',f'- 故事：{t["story"]}',f'- 引书：{t["bridge"]}',f'- CTA：{t["cta"]}',f'- 视觉：{t["visual"]}',f'- 新提议题：{t["topic"]}',f'- 参考：{", ".join(t["refs"])}',f'- 风险：{t["risk"]}',f'- 使用边界：{t["note"]}','']
atomic_text(paths['整理版模板库'],'\n'.join(tl))
readme=R/'README.md';body=readme.read_text(encoding='utf-8').split('\n## Phase 4 当前最终交付')[0]
body+='\n## Phase 4 当前最终交付\n\n全项目终验PARTIAL：25条A分析全部整合，20条限定范围AI_OK、5条AI_REVIEW；307条按计划未做深度。5套模板中4套研究可用、1套草案待补证。以reports/最终验收报告.md及manifest/phase4_final_acceptance.json为准。\n\n后续重导出使用phase4_export.py；Phase2导出命令属于历史阶段，不应覆盖当前整合版。phase4_import.py的固定视觉/更正记录适用于本批次，A补证更新后须先人工核对新增证据再刷新，不得把旧验收结论机械套用。\n'
atomic_text(readme,body)
print(json.dumps(result,ensure_ascii=False,indent=2))
