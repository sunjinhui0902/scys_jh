"""Freeze the basic Phase 2 delivery after actual data and visual reviews."""
import collections, json
from pathlib import Path
from PIL import Image
from openpyxl import load_workbook
from pipeline import OUTPUT_DIR as R, atomic_json, atomic_text, now, sha256

def read(p):return json.loads(p.read_text(encoding='utf-8'))
rows=[json.loads(s)for s in (R/'manifest/manifest.jsonl').read_text(encoding='utf-8').splitlines()if s]
lookup={x['video_id']:x for x in rows};ids=set(lookup)
plan=read(R/'manifest/phase2_plan.json');audit=read(R/'manifest/phase2_validation.json');summary=read(R/'manifest/phase2_analysis_summary.json');candidates=read(R/'clusters/deep_candidates.json')
assert len(ids)==len(rows)==332 and len(set(plan['canonical_by_video_id'].values()))==317
assert audit['status']=='PASS' and audit['source_hash_unchanged']==332 and not audit['failures']
assert audit['analysis_summary']==summary
assert audit['ready_packages']==332 and audit['verified_jpegs']==1328 and audit['resume_test_passed']
assert read(R/'state/phase2_job.json')['status']=='COMPLETE'
assert read(R/'manifest/similarity_pair_review.json')['final_pairs_reviewed']==21
assert len(candidates)==len({c['sha256']for c in candidates})==25
assert min(collections.Counter(c['account']for c in candidates).values())>=4
assert all(c['duration_s']<=1500 for c in candidates)
for c in candidates:
    v=c['video_id'];assert v in ids
    p=R/'handoff'/v;ip=read(p/'input.json');assert read(p/'READY.json')['video_id']==ip['video_id']==v
    assert ip['source_video_path']==lookup[v]['source_path']
    assert Path(ip['transcript_path']).read_text(encoding='utf-8')==(R/'transcripts'/f'{v}.txt').read_text(encoding='utf-8')
    assert len(ip['frames'])==4 and all(Path(f['path']).exists()for f in ip['frames'])
ready=read(R/'handoff/deep_candidates.READY.json');assert set(ready['video_ids'])=={c['video_id']for c in candidates}
workbook=R/'阅见生活_竞品视频分析_终版.xlsx';wb=load_workbook(workbook,data_only=False)
assert len(wb.sheetnames)==7 and len(list(R.glob('*.xlsx')))==1
for ws in wb.worksheets[:3]:
    assert ws.max_row==337 and {ws.cell(i,2).value for i in range(6,338)}==ids
for i,x in enumerate(rows,6):
    assert wb.worksheets[0].cell(i,17).value==str((x.get('db')or{}).get('video_id','未关联'))
    assert wb.worksheets[0].cell(i,17).data_type=='s'
    assert wb.worksheets[0].cell(i,22).value=='AI_PENDING'
    for j in [5,6,7,8,9,10,11,12,16,19,20,21]:assert wb.worksheets[1].cell(i,j).value=='AI_PENDING'
    for j in [11,12,13,14]:assert wb.worksheets[2].cell(i,j).value=='AI_PENDING'
assert all(wb.worksheets[4].cell(6,j).value=='AI_PENDING'for j in range(1,9))
visual=read(R/'manifest/workbook_visual_qa.json');pngcheck=read(R/'manifest/render_run_result.json');assert pngcheck['status']=='PASS'
assert len({o['sheet']for o in visual['outputs']})==7
for p in visual['outputs']:
    with Image.open(p['path'])as im:im.verify()
limitations=['Artifact render process returned exit code 1 after completing all awaited renders and PNG validation; all eight PNGs independently decoded with Pillow and actually viewed. Cause of shutdown code not established.','Artifact preview coerces long numeric platform ID strings to scientific notation; actual XLSX values are exact string cells with text format, independently compared to manifest. Preview is not the ID evidence source.']
visual.update({'reviewed_at':now(),'reviewer':'Codex','status':'PASS','sheets_visually_reviewed':7,'additional_candidate_reason_preview_reviewed':True,'observations':'Readable Chinese headers, appropriate summary wrapping, evidence links, visible pending status; candidate reason detail inspected. Original source platform IDs validated in XLSX directly.','render_process_exit_code':1,'render_engine_limitations':limitations})
atomic_json(R/'manifest/workbook_visual_qa.json',visual)
validation=read(R/'manifest/workbook_validation.json');validation.update({'status':'PASS','visual_qa_reference':str(R/'manifest/workbook_visual_qa.json'),'platform_ids_exact_string_compare':'PASS','unfinished_ai_fields_pending':'PASS','workbook_sha256':sha256(workbook)})
atomic_json(R/'manifest/workbook_validation.json',validation)
result={'at':now(),'status':'PASS','scope':'Phase 2 local baseline and seven-sheet basic workbook only','project_status':'PARTIAL','source_files':332,'sha_unique_contents':317,'asr_generated_or_sha_reused_paths':332,'verified_jpegs':1328,'ready_packages':332,'candidate_count':25,'candidate_accounts':summary['candidate_accounts'],'candidate_total_hours':sum(c['duration_s']for c in candidates)/3600,'ai_ok':0,'ai_review':3,'unfinished_ai_fields':'AI_PENDING','text_pair_review_count':21,'word_accuracy_not_claimed':True,'source_integrity_reference':str(R/'manifest/phase2_source_integrity.json'),'workbook_reference':str(workbook),'workbook_sha256':validation['workbook_sha256'],'workbook_hyperlinks_checked':validation['hyperlinks'],'asr_tail_review_pending':['v_2fe34883cba71ae1'],'render_engine_limitations':limitations}
atomic_json(R/'manifest/phase2_delivery_acceptance.json',result)
report=R/'reports/Phase2_全量本地处理验收.md'
text=report.read_text(encoding='utf-8').split('\n## 基础版交付复核')[0]
text+='\n## 基础版交付复核\n\n'
text+=f'交付时间：{result["at"]}。**Phase 2 基础处理与基础Excel交付：PASS；全项目：PARTIAL（AI深度分析与模板未完成）。**\n\n'
text+=f'单个工作簿恰好7张表，前三张各332条来源；全部video_id和长平台ID逐项核对。{validation["hyperlinks"]}个本地证据链接可达，冻结与筛选启用；未完成业务AI字段均为AI_PENDING。7张表及额外候选原因区域均已实际查看。\n\n'
text+='25条候选每账号至少4条，锦绣5条；共约2.90小时素材，均有READY包、原文与时间戳。候选清单、原因与A后续任务指令已落盘。\n\n'
text+='渲染工具局限：完成8张预览且PNG和工作簿独立读取验证通过后，工具进程仍返回非零码；原因未确认。其局部预览会把长数字字符串显示为科学计数，但Excel原文件中317个已关联平台ID均为精确文本，已逐项比对；不可从预览抄录平台ID。此局限及程序/视觉复核证据保存在manifest。\n'
atomic_text(report,text)
readme=R/'README.md';body=readme.read_text(encoding='utf-8').split('\n## Phase 2 已交付')[0]
body+='\n## Phase 2 已交付\n\n332个来源路径、317个SHA内容，全部已完成转录或同SHA复用、时间戳、4帧及READY包；一个7-Sheet基础Excel已完成。正式深度分析限定clusters/deep_candidates.json中25条，不能把332个READY包解释为全部深度分析任务。\n\n'
body+='请把reports/Antigravity_Phase3_25条任务指令.md交给A；候选共约2.90小时，每批建议5条。A只写analysis_a/，先完成临时文件并校验合法JSON后原子重命名。现有3份试标注尚待修正，核心AI字段保持AI_PENDING。\n\n'
body+='真实断点与失败恢复证据见manifest/resume_test_result.json、logs/failures/。本批相似度已逐对复核21对，固定记录脚本不应作为新素材的自动人工验收。最终阶段性证据见manifest/phase2_delivery_acceptance.json。渲染工具局限见manifest/workbook_visual_qa.json；平台ID应从Excel文本单元格或manifest读取。\n'
atomic_text(readme,body)
print(json.dumps(result,ensure_ascii=False,indent=2))
