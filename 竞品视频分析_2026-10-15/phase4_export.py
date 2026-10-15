import collections, copy, hashlib, json, re
from pathlib import Path
from urllib.parse import urlparse,unquote
from openpyxl import load_workbook
from openpyxl.styles import Font,Alignment,PatternFill
from openpyxl.worksheet.table import TableColumn
from pipeline import OUTPUT_DIR as R, atomic_json, atomic_text, now, sha256
from phase2_export import table,link,NAMES,safe,compact

def read(p):return json.loads(p.read_text(encoding='utf-8'))
def short(value,n=110):
    t=json.dumps(value,ensure_ascii=False)if isinstance(value,(dict,list))else str(value)
    return t if len(t)<=n else t[:n]+'…（完整见证据）'
def put(ws,row,col,value):
    c=ws.cell(row,col)
    if isinstance(value,dict)and'path'in value:c.value=value['label'];c.hyperlink=Path(value['path']).as_uri();c.font=Font(name='Microsoft YaHei',size=10,color='0563C1',underline='single')
    else:c.value=safe(value);c.font=Font(name='Microsoft YaHei',size=10,color='273746')
    c.fill=PatternFill(fill_type=None)
    if isinstance(c.value,str)and(c.value.startswith('AI_PENDING')or c.value=='AI_REVIEW'):c.fill=PatternFill('solid',fgColor='FFF2CC')
    if c.value=='AI_OK':c.fill=PatternFill('solid',fgColor='E2F0D9')
    c.alignment=Alignment(vertical='top',wrap_text=True);c.number_format='@'if isinstance(c.value,str)else'0.000'
def extend(ws,extra):
    for j,(name,width)in enumerate(extra,ws.max_column+1):
        c=ws.cell(5,j,name);c._style=copy.copy(ws.cell(5,1)._style);ws.column_dimensions[c.column_letter].width=width
    from openpyxl.utils import get_column_letter
    ref=f'A5:{get_column_letter(ws.max_column)}337';ws.auto_filter.ref=ref
    for t in ws.tables.values():t.ref=ref;t.tableColumns=[TableColumn(id=j,name=str(ws.cell(5,j).value))for j in range(1,ws.max_column+1)];t.autoFilter.ref=ref
def pathlink(c):
    if not c.hyperlink:return c.value
    value=unquote(urlparse(c.hyperlink.target).path)
    if re.match(r'^/[A-Za-z]:/',value):value=value[1:]
    return link(value,c.value)

def main():
    out=R/'阅见生活_竞品视频分析_终版.xlsx';oldstate=R/'manifest/phase4_export_state.json'
    files=list((R/'manifest/phase4_records').glob('v_*.json'))+[R/'manifest/phase4_templates.json',R/'manifest/phase4_account_summary.json',R/'manifest/phase4_source_integrity.json',Path(__file__)]
    digest=hashlib.sha256(''.join(sha256(p)for p in sorted(files)).encode()).hexdigest()
    if oldstate.exists():
        old=read(oldstate)
        if old['input_digest']==digest and old['output_sha256']==sha256(out):print('EXPORT_RESUME_SKIP: unchanged inputs and workbook; no rewrite');return
    backup=R/'manifest/phase4_previous_workbook.bin'
    if not backup.exists():backup.write_bytes(out.read_bytes())
    # Always rebuild from the immutable Phase 2 workbook to make reruns idempotent.
    from io import BytesIO
    wb=load_workbook(BytesIO(backup.read_bytes()));rows=[json.loads(s)for s in(R/'manifest/manifest.jsonl').read_text(encoding='utf-8').splitlines()if s];lookup={x['video_id']:x for x in rows}
    records={p.stem:read(p)for p in(R/'manifest/phase4_records').glob('v_*.json')};accounts=read(R/'manifest/phase4_account_summary.json')['accounts'];templates=read(R/'manifest/phase4_templates.json')['templates'];summary=read(R/'manifest/phase4_import_summary.json');local=read(R/'manifest/phase2_validation.json');integrity=read(R/'manifest/phase4_source_integrity.json')
    ws1,ws2,ws3=wb.worksheets[:3]
    extend(ws1,[('AI分类',28),('A声明输入模式',23),('核验范围',65),('导入证据/更正',20)])
    extend(ws2,[('A原0～3秒摘录（未经采用）',48),('A完整Hook（不等于0～3秒）',52),('Bridge终点(秒)',16),('CTA终点(秒)',16),('Bridge类型',40),('CTA类型',40),('A结尾行为原述（非必然促单）',60),('提及书名（候选另标）',38),('视频声称的产品价值（未验证效果）',60),('A推断价值（推断）',55),('导入校验与原件',20),('核验范围/更正',70)])
    extend(ws3,[('A声明输入模式',22),('视觉证据记录',20),('A原画面描述（与核验结果分列）',60),('A原声音描述（未核）',60),('补帧索引',20)])
    pain={'开卷读书':'生命无常、陪伴与衰老焦虑','文汇书香苑':'家庭阅读与理解人物选择的需求','林睿读书&书在橱窗':'历史解释与文旅知识筛选需求','破局书屋':'临场表达、社交尴尬或商业阅读困惑','赵健的读书日记':'理解城市文化与技术变化的需求','锦绣千秋录':'筛选旅行信息与看懂风景的需求'}
    nativeindex=R/'reports/phase4_visual_review/native_check/index.json'
    for i,x in enumerate(rows,6):
        v=x['video_id'];r=records.get(v)
        for ws,first,last in [(ws1,24,27),(ws2,22,33),(ws3,17,21)]:
            for j in range(first,last+1):put(ws,i,j,'AI_PENDING')
        if not r:continue
        n=r['normalized'];a=r['A_original']['fields'];ev=R/'manifest/phase4_records'/f'{v}.json';b=n['bridge'];c=n['cta'];value=n['product_value'];vis=n['visual_and_audio']
        changes={7:r['status'],22:short(n['theme']),24:n['category'],25:r['input_mode_declared'],26:r['acceptance_scope'],27:link(ev,'逐字段校验与更正')}
        for j,z in changes.items():put(ws1,i,j,z)
        changes={3:r['status'],4:n['hook_text'],5:n['hook_type'],6:short(n['theme']),7:short(n['narrative_structure']),8:n['recommended_book'],9:short(b['text']),10:b['start_s']if b['present']else'无',11:short(c['text']),12:c['start_s']if c['present']else'无',16:short(n['target_audience']),18:link(r['source_A_path'],'A原始JSON'),19:'推断（据A主题/受众整理）：'+pain[x['account']],20:short('视频声称：'+str(value.get('explicit_claims','未提供'))+'；推断：'+str(value.get('ai_inferred_value','未提供')),150),21:r['input_mode_declared'],22:short(a['hook'].get('text_0_3s','未提供')),23:short(a['hook'].get('full_hook_text','未单列')),24:b['end_s']if b['present']else'无',25:c['end_s']if c['present']else'无',26:b['type'],27:c['type'],28:short(n['closing_behavior_A']),29:'；'.join(n['mentioned_books'])or'无明确书名提及',30:short(value.get('explicit_claims','未提供')),31:'推断：'+short(value.get('ai_inferred_value','未提供')),32:link(ev,'完整字段/校验'),33:short('；'.join(r['corrections']+r['unresolved_issues']),150)}
        for j,z in changes.items():put(ws2,i,j,z)
        va=a['visual_and_audio'];changes={11:vis['visual_type'],12:vis['subtitle_style'],13:vis['book_appearance'],14:vis['audio_style'],15:r['status'],16:vis['review_method'],17:r['input_mode_declared'],18:link(ev,'四帧与核验记录'),19:short(va.get('visual_type','未提供')+'；'+va.get('book_appearance','未提供')),20:short(va.get('audio_style','未提供')),21:link(nativeindex,'指定时刻补帧')if v in['v_7e9d12d17d945cb4','v_30cc827b5d5dd936']else'无额外补帧，已看4帧'}
        for j,z in changes.items():put(ws3,i,j,z)
        ws1.row_dimensions[i].height=94;ws2.row_dimensions[i].height=112;ws3.row_dimensions[i].height=88
    ws1['A3']='Phase 4 导入版｜25条分析已整合，逐条验收状态见AI状态；未分析307条仍AI_PENDING。'
    ws2['A3']='原话取本地时间证据，A原摘录另列；AI_OK限文字/所见四帧，未知字段不冒充已识别。'
    ws3['A3']='四帧逐条核对；未验证配音/音乐明确标注。原A描述与实际所见分列。'
    freq=[]
    for a in accounts:
        source=[x for x in rows if x['account']==a['account']];topics='；'.join(a['A_original_summary']['core_topics'])
        freq.append(['账号核验汇总',a['account'],len(source),len({x['sha256']for x in source}),a['sample_count'],a['ai_ok'],a['ai_review'],a['book_introduction_count'],a['shopping_guidance_count'],short(topics+'；'+a['checked_visual_style'],145),compact(a['sample_ids']),'样本结构不可外推全账号；有引导不等于有销量',link(R/'manifest/phase4_account_summary.json','统计与逐条证据')])
    for a in accounts:
        freq.append(['A原商业分类（保留供对照）',a['account'],a['sample_count'],a['sample_count'],'不适用','不适用','不适用','见首段核验统计','见首段核验统计',a['A_original_summary']['commercial_rate_in_sample'],compact(a['sample_ids']),'原商业率混合提及/阅读/借阅/促单，未直接采用',link(R/'analysis_a/account_summary.json','A原账号总结')])
        freq.append(['A原风格判断（未证明效果）',a['account'],'不适用','不适用',a['sample_count'],'不适用','不适用','不适用','不适用',a['A_original_summary']['style'],compact(a['sample_ids']),'静态所见以首段修正版为准；不能据此认定受众或商业效果',link(R/'analysis_a/account_summary.json','A原总结')])
    for key,filename,termkey in [('全量书名候选','book_frequencies.json','term'),('全量人物/故事','story_entity_frequencies.json','term'),('全量文本关键词候选','keyword_frequencies.json','keyword')]:
        for z in read(R/'clusters'/filename)[:(30 if termkey=='keyword' else 1000)]:
            freq.append([key,z[termkey],z['video_count'],317,'非深度统计','不适用','不适用','不适用','不适用',f'涉及{z["account_count"]}账号'+(f'；标题主讲{z["title_focus_count"]}条'if'title_focus_count'in z else''),compact(e['video_id']for e in z['evidence']),'频次以SHA唯一内容为单位；提及不等于推荐',link(R/'clusters'/filename,'逐条证据')])
    for category,counts in [('Bridge类型',collections.Counter(r['normalized']['bridge']['type']for r in records.values()if r['normalized']['bridge']['present'])),('CTA类型',collections.Counter(r['normalized']['cta']['type']for r in records.values()if r['normalized']['cta']['present']))]:
        for k,n in counts.items():freq.append([category,k,n,25,25,'不适用','不适用','不适用','不适用','类型标签为A解释/整理，需结合原话','详见25条','仅深度样本；不作为高转化证据',link(R/'manifest/phase4_account_summary.json','分账号来源')])
    for filename in ['exact_duplicates.json','text_clusters.json']:
        for z in read(R/'clusters'/filename):freq.append([z['kind'],z['cluster_id'],len(z['video_ids']),317,'不适用','不适用','不适用','不适用','不适用',short(z.get('sample_quote',''),80),compact(z['video_ids']),f'涉及{z["account_count"]}账号；簇不等于任意两条都相似',link(R/'clusters'/filename,'簇与配对证据')])
    wb.remove(wb[NAMES[3]]);ws=table(wb,NAMES[3],['类型','账号/词项/簇','来源数/频次','唯一内容/分母','深度样本','AI_OK','AI_REVIEW','引书阅读数','明确商品引导数','主题/形式/类型差异','证据video_id','统计口径与局限','证据'],freq,[28,32,17,18,16,13,15,18,20,72,50,62,22],'账号总量按332来源/317SHA；深度规律仅25条样本。A原分类保留，纠正后统计另列。')
    for i in range(6,12):ws.row_dimensions[i].height=96
    tr=[]
    for t in templates:tr.append([t['id'],t['status'],t['name'],t['audience'],t['books'],t['seconds'],t['hook'],t['story'],t['bridge'],t['cta'],t['visual'],t['topic'],'；'.join(t['refs']),link(R/'manifest/phase4_templates.json','逐模板证据'),t['risk'],t['note'],link(R/'analysis_a/原创带货模板库.md','A原模板全文')])
    wb.remove(wb[NAMES[4]]);ws=table(wb,NAMES[4],['模板ID','验收状态','结构名称','适用受众（推断）','书籍范围','建议时长','开头方法','故事方法','自然引书方法','CTA方法','视觉建议','建议原创选题（待测试）','参考video_id','证据','相似性/版权风险','补证与使用边界','A原模板'],tr,[14,27,34,40,32,28,44,48,48,48,45,48,38,20,55,55,20],'5套A结构全部导入并抽象改写；4套研究结构可用、1套单样本草案。不是可直接发布的成篇文案。')
    for i in range(6,11):ws.row_dimensions[i].height=126
    oldex=wb[NAMES[5]];ex=[]
    for rr in oldex.iter_rows(min_row=6):
        if rr[2].value!='A结果校验':ex.append([pathlink(c)for c in rr])
    for v,r in records.items():
        for issue in r['schema_errors']+r['unresolved_issues']:ex.append([v,r['account'],'A未决问题',issue,r['status'],link(R/'manifest/phase4_records'/f'{v}.json','核验与原件')])
        if r['corrections']:ex.append([v,r['account'],'已整理更正',short('；'.join(r['corrections']),140),'已更正导入版，原A文件未改',link(R/'manifest/phase4_records'/f'{v}.json','更正明细')])
    ex += [['T03','模板','第二参考缺映射','“第1期山西古建”未给video_id，不能计为第二有效来源','AI_REVIEW',link(R/'analysis_a/原创带货模板库.md','A原模板')],['25条','全部','原声音判断边界','文字+四帧不支持完整配音/音乐判断，原A描述另列，核验栏不冒充已识别','已按范围限定',link(R/'manifest/phase4_import_summary.json','范围与计数')]]
    wb.remove(oldex);table(wb,NAMES[5],['video_id/模板','账号','问题类型','问题','状态/处理方式','证据'],ex,[25,22,27,82,44,23],'保留历史失败、词名歧义、更正及未决项。原A自报AI_OK不等于本项目终验通过。')
    metrics=[['全项目状态','PARTIAL','20条限定范围AI_OK；5条AI_REVIEW，原音/模态/事实补证未齐'],['导出时间',now(),'本地时区+08:00'],['来源视频文件',332,'6账号、一来源一行'],['SHA唯一内容',317,'15份完全重复额外文件仍保留来源'],['本地文本/时间戳成功路径',332,'317唯一内容识别+同SHA复用；不等于逐字准确率'],['基础图片',1328,'每来源4帧；补帧另列'],['READY交接包',332,'候选25条均有READY'],['A单视频JSON',25,'25/25与候选ID完全对齐，结构均通过'],['AI_OK',20,'限核验后的文字、时间与4帧；不代表史实/效果或全片音画均被核实'],['AI_REVIEW',5,'3条native_video证据边界；2条未核实逝世叙事'],['AI_PENDING',307,'未做深度AI分析；仍保留全部基础结果'],['A账号总结',6,'商业分类重算为引书/阅读、明确商品引导；原总结保留'],['模板数量',5,'4研究结构可用、1单样本待补证；均非发布成稿'],['全文/开头/段落簇','9 / 5 / 10','不同统计可重叠，不相加为唯一簇数'],['完全重复簇',15,'SHA一致，来源未吞并'],['跨账号相似配对',0,'当前文本阈值未命中；同书同题另计'],['配对文本复核',21,'20固定种子随机+1覆盖；Codex文本复核，非用户逐句耳听'],['本轮断点续跑','PASS','导入进程真实异常退出75，恢复后已完成5个检查点哈希/mtime不变'],['既有ASR断点复核','20/20','以前真实中断留下的已完成分段仍未重写'],['源视频不变性',integrity['unchanged_source_files'],'全量SHA、大小及mtime复核'],['A原件不变性',str(integrity['A_unchanged'])+'/28','单视频、总结、模板及能力报告只读'],['数据库','本轮未打开','元数据与Phase2快照一致'],['收费API调用',0,'本轮本地整理，无素材上传；A既有额度/费用未知'],['Phase2处理墙钟分钟',round(local['run_wall_seconds_including_restart_test']/60,2),'包括中断与补跑，不含本轮整合'],['本地工具','Python / openpyxl / FFmpeg','沿用既有环境；未安装视频框架'],['长平台ID','逐项精确文本核对','不从预览科学计数文本抄录ID'],['证据关系','A原件→冻结快照→校验记录→工作簿','字段更正可追溯，不覆盖A原文件']]
    wb.remove(wb[NAMES[6]]);table(wb,NAMES[6],['指标','实测/状态','定义与局限'],metrics,[37,46,100],'最终阶段性口径：基础332条、深度25条、核验通过20条。PARTIAL不冒充全量深度分析完成。')
    wb._sheets=[wb[name]for name in NAMES]
    tmp=out.with_suffix('.tmp.xlsx');wb.save(tmp);tmp.replace(out)
    check=load_workbook(out,data_only=False);assert check.sheetnames==NAMES
    for ws in check.worksheets[:3]:assert ws.max_row==337 and {ws.cell(i,2).value for i in range(6,338)}==set(lookup)
    states=collections.Counter(check.worksheets[0].cell(i,7).value for i in range(6,338));assert states=={'AI_OK':20,'AI_REVIEW':5,'AI_PENDING':307}
    links=0
    for ws in check:
        assert ws.freeze_panes and ws.auto_filter.ref
        for rr in ws:
            for c in rr:
                assert c.data_type!='f','Unexpected external formula'
                if c.hyperlink:
                    u=urlparse(c.hyperlink.target);p=unquote(u.path)
                    if re.match(r'^/[A-Za-z]:/',p):p=p[1:]
                    assert u.scheme=='file'and Path(p).exists(),c.hyperlink.target;links+=1
    for i,x in enumerate(rows,6):assert check.worksheets[0].cell(i,17).value==str((x.get('db')or{}).get('video_id','未关联'))and check.worksheets[0].cell(i,17).data_type=='s'
    state={'at':now(),'input_digest':digest,'output_sha256':sha256(out),'workbook':str(out),'sheets':NAMES,'rows':332,'ai_status_counts':dict(states),'hyperlinks':links,'formula_injection_scan':'PASS','long_platform_ids_exact_text':'PASS','status':'STRUCTURAL_PASS_VISUAL_PENDING'}
    atomic_json(oldstate,state)
    before=R/'manifest/phase4_manifest_before.jsonl'
    if not before.exists():before.write_bytes((R/'manifest/manifest.jsonl').read_bytes())
    for x in rows:x['ai_status']=records.get(x['video_id'],{}).get('status','AI_PENDING');x['ai_phase4_record']=str(R/'manifest/phase4_records'/f'{x["video_id"]}.json')if x['video_id']in records else None
    atomic_text(R/'manifest/manifest.jsonl',''.join(json.dumps(x,ensure_ascii=False)+'\n'for x in rows))
    print(json.dumps(state,ensure_ascii=False,indent=2))
if __name__=='__main__':main()

